import os
from pathlib import Path
import openpyxl
from .selectors import EXCEL_COLUMNS_CONFIG

class ExcelValidator:
    MANDATORY_KEYS = [
        "url"
    ]

    @staticmethod
    def is_file_locked(file_path):
        """Memeriksa apakah file Excel sedang terbuka di aplikasi lain (Microsoft Excel) yang mengunci penulisan."""
        if not os.path.exists(file_path):
            return False
        try:
            with open(file_path, "r+b"):
                return False
        except (IOError, PermissionError):
            return True

    @classmethod
    def validate_excel_file(cls, file_path):
        """
        Melakukan validasi menyeluruh terhadap file Excel template DC Usaha #3603.
        Mengembalikan dict hasil pemeriksaan.
        """
        p = Path(file_path)
        if not p.exists():
            return {
                "valid": False,
                "error": f"File tidak ditemukan: {file_path}",
                "warning": None,
                "file_locked": False,
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        locked = cls.is_file_locked(file_path)
        warning_msg = None
        if locked:
            warning_msg = "Perhatian: File sedang dibuka di Excel. Anda dapat melihat pratinjau, namun harap tutup file Excel sebelum memulai otomasi agar hasil dapat disimpan."

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheet = wb.active
        except Exception as e:
            return {
                "valid": False,
                "error": f"Gagal membaca file Excel: {e}",
                "warning": None,
                "file_locked": locked,
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        # Baca header baris 1
        header_map = {}
        for col_idx in range(1, sheet.max_column + 1):
            val = sheet.cell(row=1, column=col_idx).value
            if val is not None:
                header_map[str(val).strip().lower()] = col_idx

        # Mapping konfigurasi ke index kolom sebenarnya
        col_map = {}
        for key, aliases in EXCEL_COLUMNS_CONFIG.items():
            for alias in aliases:
                clean_alias = str(alias).strip().lower()
                if clean_alias in header_map:
                    col_map[key] = header_map[clean_alias]
                    break

        # Periksa kolom wajib (hanya url)
        missing_mandatory = []
        for mand in cls.MANDATORY_KEYS:
            if mand not in col_map:
                missing_mandatory.append(mand.replace("_", " ").title())

        if missing_mandatory:
            wb.close()
            return {
                "valid": False,
                "error": f"Template tidak sesuai. Kolom wajib berikut tidak ditemukan: {', '.join(missing_mandatory)}",
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        # Parsing baris data untuk preview & counting
        rows_data = []
        total_rows = 0
        done_rows = 0
        pending_rows = 0

        url_col = col_map.get("url")
        name_col = col_map.get("nama_usaha")
        id_col = col_map.get("id")
        status_col = col_map.get("status")
        ket_col = col_map.get("keterangan")

        for r in range(2, sheet.max_row + 1):
            url_val = sheet.cell(row=r, column=url_col).value if url_col else None
            if url_val is None or str(url_val).strip() == "":
                continue

            total_rows += 1
            st_val = str(sheet.cell(row=r, column=status_col).value or "").strip() if status_col else ""
            ket_val = str(sheet.cell(row=r, column=ket_col).value or "").strip() if ket_col else ""
            nama_val = str(sheet.cell(row=r, column=name_col).value or "").strip() if name_col else ""
            asg_id = str(sheet.cell(row=r, column=id_col).value or "").strip() if id_col else str(r - 1)

            is_done = st_val.upper() in ["SUKSES", "DONE", "SELESAI", "BERHASIL"]
            if is_done:
                done_rows += 1
            else:
                pending_rows += 1

            rows_data.append({
                "row_idx": r,
                "id": asg_id,
                "nama_usaha": nama_val,
                "url": str(url_val).strip(),
                "status": st_val if st_val else ("SUKSES" if is_done else "MENUNGGU"),
                "keterangan": ket_val
            })

        wb.close()

        return {
            "valid": True,
            "error": None,
            "warning": warning_msg,
            "file_locked": locked,
            "total_rows": total_rows,
            "done_rows": done_rows,
            "pending_rows": pending_rows,
            "rows_data": rows_data
        }

    @staticmethod
    def generate_template(output_path):
        """Membuat file template Excel awal untuk DC Usaha #3603."""
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DC_Usaha_3603"

        headers = [
            "No", "ID Penugasan", "Nama Usaha", "Link Fasih SM", "Status", "Keterangan"
        ]

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        sample_rows = [
            [
                1, "3603-DC-001", "CV Contoh Berkah Mandiri",
                "https://fasih-sm.bps.go.id/penugasan/detail/sample-dc-1",
                None, None
            ],
            [
                2, "3603-DC-002", "Toko Sejahtera Bersama",
                "https://fasih-sm.bps.go.id/penugasan/detail/sample-dc-2",
                None, None
            ]
        ]

        data_font = Font(name="Calibri", size=10)
        for row_data in sample_rows:
            ws.append(row_data)
            row_idx = ws.max_row
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = thin_border

        # Auto fit column width
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val_str = str(cell.value or "")
                if "\n" in val_str:
                    val_str = max(val_str.split("\n"), key=len)
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        wb.save(output_path)
        wb.close()
        return True

    generate_template_excel = generate_template

DCUsaha3603Validator = ExcelValidator
