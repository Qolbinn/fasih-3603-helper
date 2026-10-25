import os
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from modules.cleaning_nik_r30.selectors import EXCEL_NIK_COLUMNS_CONFIG

class ExcelNIKValidator:
    MANDATORY_KEYS = [
        "assignment_id",
        "no_urut_kk",
        "nomor_induk_kependudukan",
        "url",
        "nomor_kartu_keluarga",
        "nama"
    ]

    @staticmethod
    def is_file_locked(file_path):
        """Memeriksa apakah file Excel sedang terbuka di aplikasi lain (Microsoft Excel) yang mengunci penulisan."""
        if not os.path.exists(file_path):
            return False
        try:
            with open(file_path, "r+b"):
                return False
        except (IOError, PermissionError):
            return True

    @classmethod
    def validate_excel_file(cls, file_path):
        """
        Melakukan validasi menyeluruh terhadap file Excel input Data Cleaning NIK.
        Mengembalikan dict hasil pemeriksaan & data pratinjau.
        """
        p = Path(file_path)
        if not p.exists():
            return {
                "valid": False,
                "error": f"File tidak ditemukan: {file_path}",
                "warning": None,
                "file_locked": False,
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        locked = cls.is_file_locked(file_path)
        warning_msg = None
        if locked:
            warning_msg = "Perhatian: File sedang dibuka di Excel. Anda dapat melihat pratinjau, namun harap tutup file Excel sebelum memulai otomasi agar hasil dapat disimpan."

        try:
            with open(file_path, "rb") as f:
                wb = openpyxl.load_workbook(f, data_only=True)
                sheet = wb.active
        except Exception as e:
            return {
                "valid": False,
                "error": f"Gagal membaca format file Excel: {e}",
                "warning": None,
                "file_locked": locked,
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        # 1. Baca header baris 1
        header_map = {}
        for col_idx in range(1, sheet.max_column + 1):
            val = sheet.cell(row=1, column=col_idx).value
            if val is not None:
                header_map[str(val).strip().lower()] = col_idx

        # 2. Mapping konfigurasi ke index kolom sebenarnya
        col_map = {}
        for key, aliases in EXCEL_NIK_COLUMNS_CONFIG.items():
            for alias in aliases:
                clean_alias = str(alias).strip().lower()
                if clean_alias in header_map:
                    col_map[key] = header_map[clean_alias]
                    break

        # 3. Periksa kolom wajib
        missing_mandatory = []
        for mand in cls.MANDATORY_KEYS:
            if mand not in col_map:
                missing_mandatory.append(mand)

        if missing_mandatory:
            wb.close()
            return {
                "valid": False,
                "error": f"Template tidak sesuai. Kolom wajib berikut tidak ditemukan: {', '.join(missing_mandatory)}",
                "warning": warning_msg,
                "file_locked": locked,
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        # 4. Pindai baris data
        assignment_col = col_map.get("assignment_id")
        no_urut_col = col_map.get("no_urut_kk")
        nik_col = col_map.get("nomor_induk_kependudukan")
        url_col = col_map.get("url")
        no_kk_col = col_map.get("nomor_kartu_keluarga")
        nama_col = col_map.get("nama")
        status_col = col_map.get("status")
        ket_col = col_map.get("keterangan")

        total_rows = 0
        done_rows = 0
        pending_rows = 0
        rows_data = []

        max_row = sheet.max_row
        for r in range(2, max_row + 1):
            url_val = sheet.cell(row=r, column=url_col).value if url_col else None
            # Jika kolom URL kosong, baris diabaikan (baris kosong di bawah data)
            if url_val is None or str(url_val).strip() == "":
                continue

            total_rows += 1

            assignment_val = sheet.cell(row=r, column=assignment_col).value if assignment_col else ""
            no_urut_val = sheet.cell(row=r, column=no_urut_col).value if no_urut_col else ""
            nik_val = sheet.cell(row=r, column=nik_col).value if nik_col else ""
            no_kk_val = sheet.cell(row=r, column=no_kk_col).value if no_kk_col else ""
            nama_val = sheet.cell(row=r, column=nama_col).value if nama_col else ""
            st_val = sheet.cell(row=r, column=status_col).value if status_col else ""
            ket_val = sheet.cell(row=r, column=ket_col).value if ket_col else ""

            st_clean = str(st_val or "").strip().upper()
            if "SUKSES" in st_clean or "DONE" in st_clean:
                done_rows += 1
            else:
                pending_rows += 1

            # Format NIK & No KK agar tidak scientific notation jika numeric
            def clean_id_str(val):
                if val is None:
                    return ""
                s = str(val).strip()
                if s.endswith(".0"):
                    s = s[:-2]
                return s

            rows_data.append({
                "row_idx": r,
                "assignment_id": clean_id_str(assignment_val),
                "no_urut_kk": clean_id_str(no_urut_val),
                "nomor_induk_kependudukan": clean_id_str(nik_val),
                "url": str(url_val).strip(),
                "nomor_kartu_keluarga": clean_id_str(no_kk_val),
                "nama": str(nama_val or "").strip(),
                "status": str(st_val or "").strip(),
                "keterangan": str(ket_val or "").strip()
            })

        wb.close()

        if total_rows == 0:
            return {
                "valid": False,
                "error": "File Excel valid secara header, namun tidak ada baris data penugasan (kolom URL kosong di semua baris).",
                "warning": warning_msg,
                "file_locked": locked,
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        return {
            "valid": True,
            "error": None,
            "warning": warning_msg,
            "file_locked": locked,
            "total_rows": total_rows,
            "done_rows": done_rows,
            "pending_rows": pending_rows,
            "rows_data": rows_data,
            "col_map": col_map
        }

    @classmethod
    def generate_template_excel(cls, output_path):
        """
        Membuat template Excel standar untuk modul Data Cleaning Sosial Ekonomi (NIK).
        Berisi header kolom wajib dan beberapa baris contoh data.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data_Cleaning_NIK"

        headers = [
            "assignment_id",
            "no_urut_kk",
            "nomor_induk_kependudukan",
            "url",
            "nomor_kartu_keluarga",
            "nama",
            "Status_Eksekusi",
            "Keterangan"
        ]

        ws.append(headers)

        # Style header
        header_fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")  # Teal aesthetic
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='D1D5DB'),
            right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'),
            bottom=Side(style='thin', color='D1D5DB')
        )

        ws.row_dimensions[1].height = 28

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border

        # Baris Contoh Data
        sample_rows = [
            [
                "360301000100101",
                "1",
                "3603010101900001",
                "https://fasih-sm.bps.go.id/penugasan/detail/contoh-id-1",
                "3603010101200001",
                "AHMAD SUDIRMAN",
                None,
                None
            ],
            [
                "360301000100102",
                "2",
                "3603014101920002",
                "https://fasih-sm.bps.go.id/penugasan/detail/contoh-id-2",
                "3603010101200001",
                "SITI AMINAH",
                None,
                None
            ]
        ]

        for row_data in sample_rows:
            ws.append(row_data)

        # Format border & alignment pada data baris contoh
        for r_idx in range(2, len(sample_rows) + 2):
            ws.row_dimensions[r_idx].height = 22
            for c_idx in range(1, len(headers) + 1):
                c = ws.cell(row=r_idx, column=c_idx)
                c.border = thin_border
                c.font = Font(name="Segoe UI", size=10)
                # Text align center untuk ID, no_urut, NIK, No KK, status
                if c_idx in [1, 2, 3, 5, 7]:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

        # Auto width columns
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        wb.save(output_path)
        return True

ExcelNIKR30Validator = ExcelNIKValidator
