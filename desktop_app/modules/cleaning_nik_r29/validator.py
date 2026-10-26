import os
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from modules.cleaning_nik_r29.selectors import EXCEL_NIK_R29_COLUMNS_CONFIG

class ExcelNIKR29Validator:
    MANDATORY_KEYS = [
        "no_urut_kk",
        "nik_dtsen",
        "nik_length",
        "matched",
        "nama_dtsen",
        "url"
    ]

    @staticmethod
    def is_file_locked(file_path):
        """Memeriksa apakah file Excel sedang terbuka di aplikasi lain (Microsoft Excel) yang mengunci penulisan."""
        if not os.path.exists(file_path):
            return False
        try:
            with open(file_path, "r+b"):
                return False
        except (IOError, PermissionError):
            return True

    @classmethod
    def validate_excel_file(cls, file_path):
        """
        Melakukan validasi menyeluruh terhadap file Excel input Data Cleaning NIK R29.
        Mengembalikan dict hasil pemeriksaan & data pratinjau.
        """
        p = Path(file_path)
        if not p.exists():
            return {
                "valid": False,
                "error": f"File tidak ditemukan: {file_path}",
                "warning": None,
                "file_locked": False,
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        locked = cls.is_file_locked(file_path)
        warning_msg = None
        if locked:
            warning_msg = "Perhatian: File sedang dibuka di Excel. Anda dapat melihat pratinjau, namun harap tutup file Excel sebelum memulai otomasi agar hasil dapat disimpan."

        try:
            with open(file_path, "rb") as f:
                wb = openpyxl.load_workbook(f, data_only=True)
                sheet = wb.active
        except Exception as e:
            return {
                "valid": False,
                "error": f"Gagal membaca format file Excel: {e}",
                "warning": None,
                "file_locked": locked,
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        # 1. Baca header baris 1
        header_map = {}
        for col_idx in range(1, sheet.max_column + 1):
            val = sheet.cell(row=1, column=col_idx).value
            if val is not None:
                header_map[str(val).strip().lower()] = col_idx

        # 2. Mapping konfigurasi ke index kolom sebenarnya
        col_map = {}
        for key, aliases in EXCEL_NIK_R29_COLUMNS_CONFIG.items():
            for alias in aliases:
                clean_alias = str(alias).strip().lower()
                if clean_alias in header_map:
                    col_map[key] = header_map[clean_alias]
                    break

        # 3. Periksa kolom wajib
        missing_mandatory = []
        for mand in cls.MANDATORY_KEYS:
            if mand not in col_map:
                missing_mandatory.append(mand)

        if missing_mandatory:
            wb.close()
            return {
                "valid": False,
                "error": f"Template tidak sesuai. Kolom wajib berikut tidak ditemukan: {', '.join(missing_mandatory)}",
                "warning": warning_msg,
                "file_locked": locked,
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        # 4. Pindai baris data
        no_urut_col = col_map.get("no_urut_kk")
        nama_col = col_map.get("nama_dtsen")
        nik_col = col_map.get("nik_dtsen")
        len_col = col_map.get("nik_length")
        matched_col = col_map.get("matched")
        url_col = col_map.get("url")
        status_col = col_map.get("status")
        ket_col = col_map.get("keterangan")

        total_rows = 0
        done_rows = 0
        pending_rows = 0
        rows_data = []

        max_row = sheet.max_row
        for r in range(2, max_row + 1):
            url_val = sheet.cell(row=r, column=url_col).value if url_col else None
            # Jika kolom URL kosong, baris diabaikan (baris kosong di bawah data)
            if url_val is None or str(url_val).strip() == "":
                continue

            total_rows += 1

            no_urut_val = sheet.cell(row=r, column=no_urut_col).value if no_urut_col else ""
            nama_val = sheet.cell(row=r, column=nama_col).value if nama_col else ""
            nik_val = sheet.cell(row=r, column=nik_col).value if nik_col else ""
            len_val = sheet.cell(row=r, column=len_col).value if len_col else ""
            matched_val = sheet.cell(row=r, column=matched_col).value if matched_col else ""
            st_val = sheet.cell(row=r, column=status_col).value if status_col else ""
            ket_val = sheet.cell(row=r, column=ket_col).value if ket_col else ""

            st_clean = str(st_val or "").strip().upper()
            if "SUKSES" in st_clean or "DONE" in st_clean:
                done_rows += 1
            else:
                pending_rows += 1

            # Format teks angka agar tidak scientific notation jika numeric
            def clean_id_str(val):
                if val is None:
                    return ""
                s = str(val).strip()
                if s.endswith(".0"):
                    s = s[:-2]
                return s

            rows_data.append({
                "row_idx": r,
                "no_urut_kk": clean_id_str(no_urut_val),
                "nama_dtsen": str(nama_val or "").strip(),
                "nik_dtsen": clean_id_str(nik_val),
                "nik_length": clean_id_str(len_val),
                "matched": clean_id_str(matched_val),
                "url": str(url_val).strip(),
                "status": str(st_val or "").strip(),
                "keterangan": str(ket_val or "").strip()
            })

        wb.close()

        if total_rows == 0:
            return {
                "valid": False,
                "error": "File Excel valid secara header, namun tidak ada baris data penugasan (kolom URL kosong di semua baris).",
                "warning": warning_msg,
                "file_locked": locked,
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        return {
            "valid": True,
            "error": None,
            "warning": warning_msg,
            "file_locked": locked,
            "total_rows": total_rows,
            "done_rows": done_rows,
            "pending_rows": pending_rows,
            "rows_data": rows_data,
            "col_map": col_map
        }

    @classmethod
    def generate_template_excel(cls, output_path):
        """
        Membuat template Excel standar untuk modul Data Cleaning Sosial Ekonomi (NIK R29).
        Berisi header kolom wajib dan beberapa baris contoh data.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data_Cleaning_NIK_R29"

        headers = [
            "no_urut_kk",
            "nama_dtsen",
            "nik_dtsen",
            "nik_length",
            "matched",
            "url",
            "Status_Eksekusi",
            "Keterangan"
        ]

        ws.append(headers)

        # Style header
        header_fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")  # Teal aesthetic
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='D1D5DB'),
            right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'),
            bottom=Side(style='thin', color='D1D5DB')
        )

        ws.row_dimensions[1].height = 28

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border

        # Baris Contoh Data
        sample_rows = [
            [
                "1",
                "AHMAD SUDIRMAN",
                "3603010101900001",
                "16",
                None,
                "https://fasih-sm.bps.go.id/penugasan/detail/contoh-id-1",
                None,
                None
            ],
            [
                "2",
                "SITI AMINAH",
                "3603014101920002",
                "16",
                "3603014101920009",
                "https://fasih-sm.bps.go.id/penugasan/detail/contoh-id-2",
                None,
                None
            ],
            [
                "3",
                "BUDI SANTOSO",
                "36030111",
                "8",
                "7777",
                "https://fasih-sm.bps.go.id/penugasan/detail/contoh-id-3",
                None,
                None
            ],
            [
                "4",
                "RINA WATI",
                "36030122",
                "8",
                "9999",
                "https://fasih-sm.bps.go.id/penugasan/detail/contoh-id-4",
                None,
                None
            ]
        ]

        for row_data in sample_rows:
            ws.append(row_data)

        # Format border & alignment pada data baris contoh
        for r_idx in range(2, len(sample_rows) + 2):
            ws.row_dimensions[r_idx].height = 22
            for c_idx in range(1, len(headers) + 1):
                c = ws.cell(row=r_idx, column=c_idx)
                c.border = thin_border
                c.font = Font(name="Segoe UI", size=10)
                # Text align center untuk no_urut, NIK, length, matched, status
                if c_idx in [1, 3, 4, 5, 7]:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

        # Auto width columns
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        wb.save(output_path)
        return True

ExcelNIKR29Validator = ExcelNIKR29Validator


ExcelNIKValidator = ExcelNIKR29Validator
