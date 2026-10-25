import os
import re
import time
import random
import threading
import queue
import openpyxl
from datetime import datetime
from pathlib import Path
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

from config import CDP_URL
from core.logger import logger
from core.event_bus import EngineState
from core.audit_logger import audit_logger
from core.window_controller import WindowController
from core.browser_manager import BrowserManager
from modules.cleaning_nik_r29.selectors import EXCEL_NIK_R29_COLUMNS_CONFIG, NIK_R29_SELECTORS
from modules.cleaning_nik_r29.validator import ExcelNIKR29Validator

def human_jitter(min_sec=0.1, max_sec=0.3, event_bus=None):
    """Jeda acak mikro untuk mensimulasikan interaksi manusia."""
    if event_bus and event_bus.is_stop_requested():
        return
    t = random.uniform(min_sec, max_sec)
    time.sleep(t)

def random_sleep(min_sec=0.5, max_sec=1.5, event_bus=None, allow_hesitation=True):
    """Jeda acak antar aksi tombol/navigasi."""
    if event_bus and event_bus.is_stop_requested():
        return
    base = random.uniform(min_sec, max_sec)
    if allow_hesitation and random.random() < 0.15:
        base += random.uniform(0.5, 1.2)
    time.sleep(base)

def safe_click(page, selector_or_locator, timeout=20000, event_bus=None, force=False):
    """Melakukan scroll, jitter, lalu klik secara aman pada elemen."""
    if event_bus and event_bus.is_stop_requested():
        return
    if isinstance(selector_or_locator, str):
        loc = page.locator(selector_or_locator).first
    else:
        loc = selector_or_locator
    loc.wait_for(state="attached", timeout=timeout)
    try:
        loc.scroll_into_view_if_needed(timeout=5000)
    except Exception:
        pass
    human_jitter(0.1, 0.2, event_bus=event_bus)
    try:
        loc.click(timeout=timeout, force=force)
    except Exception as e:
        err_msg = str(e).lower()
        if "intercepts pointer events" in err_msg or "subtree intercepts" in err_msg or not force:
            logger.bot("Intersepsi klik terdeteksi. Mengulang klik dengan force=True...")
            human_jitter(0.1, 0.2, event_bus=event_bus)
            loc.click(timeout=8000, force=True)
        else:
            raise e

def normalize_text(text):
    """Normalisasi string: lowercase, hilangkan non-breaking space, collapse whitespace."""
    if text is None:
        return ""
    t = str(text).replace("\xa0", " ").strip().lower()
    return re.sub(r"\s+", " ", t)

def clean_nik_str(val):
    """Membersihkan nilai NIK agar murni angka digit (menghilangkan spasi, titik, strip, .0)."""
    if val is None:
        return ""
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return re.sub(r"[^\d]", "", s)

def add_remarks_to_element(page, container_locator_or_sel, note_text, event_bus=None, worker_prefix=""):
    """Membuka modal Catatan pada elemen container, mengetik catatan, klik simpan, lalu tutup modal."""
    if isinstance(container_locator_or_sel, str):
        container = page.locator(container_locator_or_sel).first
    else:
        container = container_locator_or_sel
    container.wait_for(state="attached", timeout=30000)
    try:
        container.scroll_into_view_if_needed(timeout=5000)
    except Exception:
        pass
    random_sleep(0.4, 0.7, event_bus=event_bus)

    # Cari tombol Catatan yang VISIBLE di dalam container
    btn_catatan = container.locator("button[title='Catatan']:visible, button[aria-label='Catatan']:visible").first
    if btn_catatan.count() == 0 or not btn_catatan.is_visible():
        all_btns = container.locator("button[title='Catatan'], button[aria-label='Catatan']")
        btn_catatan = None
        for i in range(all_btns.count()):
            b = all_btns.nth(i)
            if b.is_visible():
                btn_catatan = b
                break
        if not btn_catatan and all_btns.count() > 0:
            btn_catatan = all_btns.first

    if not btn_catatan:
        raise Exception("Tombol Catatan visible tidak ditemukan pada container NIK")

    safe_click(page, btn_catatan, timeout=20000, event_bus=event_bus)
    random_sleep(0.7, 1.1, event_bus=event_bus)

    dialog = page.locator("div[role='dialog']").first
    try:
        dialog.wait_for(state="visible", timeout=25000)
    except PlaywrightTimeoutError:
        raise Exception("Modal pop-up Catatan tidak muncul setelah tombol Catatan diklik")

    note_str = str(note_text).strip()

    # Cek apakah catatan sudah pernah tercatat sebelumnya di riwayat pesan
    already_exists = False
    try:
        bubbles = dialog.locator("div[class*='whitespace-pre-wrap'], div[class*='wrap-break-word']")
        for i in range(bubbles.count()):
            if bubbles.nth(i).inner_text().strip() == note_str:
                already_exists = True
                break
    except Exception:
        pass

    if already_exists:
        logger.info(f"{worker_prefix}Catatan '{note_str}' sudah tercatat sebelumnya. Melewati pengiriman ulang.")
    else:
        textarea = dialog.locator("textarea[placeholder*='Tambah catatan'], textarea").first
        try:
            textarea.wait_for(state="visible", timeout=25000)
        except PlaywrightTimeoutError:
            raise Exception("Textarea Catatan tidak ditemukan di dalam dialog")

        textarea.click()
        human_jitter(0.1, 0.2, event_bus=event_bus)
        textarea.fill(note_str)
        human_jitter(0.15, 0.25, event_bus=event_bus)

        btn_simpan = dialog.locator("button[title='Simpan']").first
        btn_simpan.wait_for(state="visible", timeout=25000)
        safe_click(page, btn_simpan, timeout=20000, event_bus=event_bus)
        random_sleep(0.8, 1.2, event_bus=event_bus)

    try:
        btn_close = dialog.locator("button[aria-label='Dismiss'], button[title='Tutup'], button:has-text('Tutup')").first
        if btn_close.is_visible(timeout=4000):
            safe_click(page, btn_close, timeout=8000, event_bus=event_bus)
        else:
            page.keyboard.press("Escape")
    except Exception:
        page.keyboard.press("Escape")

    try:
        dialog.wait_for(state="hidden", timeout=8000)
    except Exception:
        page.keyboard.press("Escape")

    random_sleep(0.4, 0.7, event_bus=event_bus)

def ensure_anomali_switch_active(page, event_bus=None, wait_timeout=35000, worker_prefix=""):
    """
    Pastikan switch 'Tampilkan Anomali Usaha dan Keluarga' (div#cek_anomali_button) AKTIF di Blok Catatan.
    """
    box = page.locator("div#cek_anomali_button").first
    try:
        box.wait_for(state="visible", timeout=wait_timeout)
    except PlaywrightTimeoutError:
        raise Exception("Switch 'Tampilkan Anomali Usaha dan Keluarga' tidak muncul di Blok Catatan.")

    try:
        box.scroll_into_view_if_needed(timeout=5000)
    except Exception:
        pass
    sw_input = box.locator("input[role='switch']").first
    sw_control = box.locator("[id^='switch-'][id$='-control']").first

    def _is_active():
        return sw_input.get_attribute("aria-checked") == "true"

    if _is_active():
        logger.bot(f"{worker_prefix}Switch 'Tampilkan Anomali Usaha dan Keluarga' sudah aktif -> Dibiarkan.")
        return False

    logger.bot(f"{worker_prefix}Switch 'Tampilkan Anomali Usaha dan Keluarga' belum aktif -> Mengaktifkan...")
    for attempt in range(1, 4):
        safe_click(page, sw_control, timeout=5000, event_bus=event_bus)
        for _ in range(10):
            random_sleep(0.25, 0.35, event_bus=event_bus, allow_hesitation=False)
            if _is_active():
                logger.bot(f"{worker_prefix}Switch 'Tampilkan Anomali Usaha dan Keluarga' berhasil diaktifkan.")
                return True
        logger.warning(f"{worker_prefix}Switch anomali belum aktif setelah klik (percobaan {attempt}/3). Mengulang...")

    raise Exception("Gagal mengaktifkan toggle Tampilkan Anomali Usaha dan Keluarga.")

def _get_valid_art_items(page):
    """
    Mengambil daftar item anggota keluarga (ART) yang valid dan clickable dari sub-sidebar Blok I.
    Mengembalikan list of tuple: (locator, clean_nama_art)
    """
    sel = (
        "xpath=//div[@title='Keterangan Anggota Keluarga']/following-sibling::div[@data-expanded]"
        "//div[contains(@class, 'fasih-form-sidebar')]/div[contains(@class, 'cursor-pointer') and @title]"
    )
    subs = page.locator(sel)
    cnt = subs.count()
    if cnt == 0:
        sel_fallback = (
            "xpath=//div[@title='Keterangan Anggota Keluarga']/following-sibling::div"
            "//div[contains(@class, 'fasih-form-sidebar')]/div[contains(@class, 'cursor-pointer') and @title]"
        )
        subs = page.locator(sel_fallback)
        cnt = subs.count()

    valid_items = []
    for idx in range(cnt):
        item = subs.nth(idx)
        raw_title = item.get_attribute("title") or item.inner_text() or ""
        clean_title = raw_title.strip()
        if not clean_title or clean_title in [
            "SE2026 - L BLOK I", "Keterangan Anggota Keluarga",
            "CATATAN", "KETERANGAN UMUM KELUARGA", "KETERANGAN PEMBERI JAWABAN"
        ]:
            continue
        valid_items.append((item, clean_title))
    return valid_items

def _find_and_open_matching_art(page, excel_no_urut, excel_nama, event_bus=None, worker_prefix="", check_stop_fn=None):
    """
    Mencari dan membuka halaman rincian ART yang cocok antara urutan (no_urut_kk) dan nama di Excel.
    Alur:
    1. Pastikan menu Blok I di sidebar diklik dan terbuka (data-expanded).
    2. Tunggu subsidebar Keterangan Anggota Keluarga terbuka otomatis.
    3. Langsung telusuri nama-nama anggota keluarga yang muncul di subsidebar.
    """
    try:
        page.locator(".fasih-form-sidebar, div[class*='sidebar']").first.wait_for(state="visible", timeout=35000)
    except Exception:
        pass

    # 1. Pastikan tombol Blok I ditemukan di sidebar
    blok1_btn = page.locator(NIK_R29_SELECTORS["sidebar_blok_i"]).first
    try:
        blok1_btn.wait_for(state="visible", timeout=30000)
    except PlaywrightTimeoutError:
        return False, None, "Menu 'SE2026 - L BLOK I' tidak ditemukan di sidebar"

    # Cek apakah Blok I sudah terbuka (data-expanded)
    is_b1_expanded = page.locator(
        "xpath=//div[contains(@class, 'fasih-form-sidebar')]//div[@title='SE2026 - L BLOK I' or contains(@title, 'BLOK I')]/following-sibling::div[@data-expanded]"
    ).count() > 0

    if not is_b1_expanded:
        logger.bot(f"{worker_prefix}Membuka menu Blok I (Keterangan Umum Keluarga) di sidebar...")
        safe_click(page, blok1_btn, timeout=20000, event_bus=event_bus)
        try:
            page.locator(
                "xpath=//div[contains(@class, 'fasih-form-sidebar')]//div[@title='SE2026 - L BLOK I' or contains(@title, 'BLOK I')]/following-sibling::div[@data-expanded]"
            ).first.wait_for(state="attached", timeout=15000)
        except Exception:
            pass
        random_sleep(0.8, 1.2, event_bus=event_bus)
    else:
        logger.bot(f"{worker_prefix}Menu Blok I sudah dalam kondisi terbuka di sidebar.")

    # 2. Tunggu subsidebar Keterangan Anggota Keluarga terbentang dan daftar ART ter-render
    valid_items = []
    for _ in range(15):
        valid_items = _get_valid_art_items(page)
        if len(valid_items) > 0:
            break
        random_sleep(0.4, 0.7, event_bus=event_bus)

    count = len(valid_items)
    if count == 0:
        return False, None, "Tidak ada rincian anggota keluarga di subsidebar Blok I"


    # Parsing no urut KK dari Excel (1-indexed)
    try:
        raw_urut = clean_nik_str(excel_no_urut)
        target_urut = int(raw_urut) if raw_urut else 1
    except Exception:
        target_urut = 1

    art_names = [name for _, name in valid_items]
    logger.bot(f"{worker_prefix}Ditemukan {count} anggota keluarga di subsidebar: {art_names}")
    logger.bot(f"{worker_prefix}Mencocokkan urutan ke-{target_urut} dengan target nama '{excel_nama}'...")

    if target_urut < 1 or target_urut > count:
        return False, None, f"No urut KK {target_urut} di luar jumlah anggota keluarga di web ({count} orang)"

    # Cek anggota keluarga pada posisi target_urut (1-based index)
    target_pos_idx = target_urut - 1
    curr_loc, curr_name = valid_items[target_pos_idx]

    # Validasi kesesuaian nama
    norm_excel_nama = normalize_text(excel_nama)
    norm_curr_name = normalize_text(curr_name)

    if norm_excel_nama == norm_curr_name:
        logger.success(f"{worker_prefix}Anggota keluarga COCOK pada urutan ke-{target_urut}: '{curr_name}'")
        safe_click(page, curr_loc, timeout=20000, event_bus=event_bus)
        random_sleep(1.2, 1.8, event_bus=event_bus)
        return True, curr_name, None
    else:
        # Cek apakah nama tersebut ada di posisi lain
        pos_found = None
        for idx, (_, name) in enumerate(valid_items):
            if normalize_text(name) == norm_excel_nama:
                pos_found = idx + 1
                break

        if pos_found:
            err_msg = f"Nama '{excel_nama}' ditemukan pada urutan ke-{pos_found}, tetapi di Excel tercatat urutan ke-{target_urut}. Urutan tidak cocok."
        else:
            err_msg = f"Anggota keluarga pada urutan ke-{target_urut} adalah '{curr_name}', tidak cocok dengan nama di Excel '{excel_nama}'."

        logger.warning(f"⚠️ {worker_prefix}{err_msg}")
        return False, curr_name, err_msg


def _verify_and_remark_nik(page, excel_nik, event_bus=None, worker_prefix=""):
    """
    Memeriksa isian pada rincian NIK (id nik_dtsen).
    - Jika SAMA dengan nomor_induk_kependudukan di Excel: berikan remarks '#3' pada catatan rincian NIK.
    - Jika BEDA: return False beserta keterangan beda.
    """
    nik_box = page.locator(NIK_R29_SELECTORS["container_nik"]).first
    try:
        nik_box.wait_for(state="attached", timeout=25000)
        nik_box.scroll_into_view_if_needed(timeout=5000)
    except Exception:
        pass

    if nik_box.count() == 0:
        raise Exception("Elemen rincian NIK (id nik_dtsen) tidak ditemukan di form")

    # Ambil nilai NIK yang terisi di web
    inp_nik = nik_box.locator(NIK_R29_SELECTORS["input_nik"]).first
    try:
        inp_nik.wait_for(state="visible", timeout=15000)
    except PlaywrightTimeoutError:
        raise Exception("Field input NIK tidak visible di halaman formulir")

    web_nik = ""
    try:
        web_nik = inp_nik.input_value().strip()
    except Exception:
        pass
    if not web_nik:
        web_nik = (inp_nik.get_attribute("value") or "").strip()

    clean_web = clean_nik_str(web_nik)
    clean_excel = clean_nik_str(excel_nik)

    logger.bot(f"{worker_prefix}Memeriksa NIK (Web: '{web_nik}' vs Excel: '{excel_nik}')...")

    if clean_web == clean_excel and clean_excel != "":
        logger.success(f"{worker_prefix}NIK COCOK! (Web: '{web_nik}' == Excel: '{excel_nik}')")
        logger.remarks(f"{worker_prefix}Memberikan remarks '#3' pada Catatan rincian NIK...")
        add_remarks_to_element(page, nik_box, "#3", event_bus=event_bus, worker_prefix=worker_prefix)
        return True, web_nik, None
    else:
        err_msg = f"NIK tidak sama (Web: '{web_nik}' vs Excel: '{excel_nik}')"
        logger.warning(f"⚠️ {worker_prefix}{err_msg}")
        return False, web_nik, err_msg

def _submit_assignment_form(page, event_bus, row_idx, nama_art, debug_review=False, auto_minimize=True, worker_prefix=""):
    """
    Menjalankan alur submit:
    1. Buka Blok Catatan
    2. Pastikan switch toggle anomali aktif
    3. Klik Kirim formulir Fasih
    4. Cek Galat pada dialog pop-up (batal jika > 0)
    5. Konfirmasi Kirim (dengan dukungan jeda review pada Mode Debugging)
    6. Verifikasi dialog sukses
    """
    # 1. Buka Blok Catatan
    logger.bot(f"{worker_prefix}Membuka menu Blok Catatan...")
    try:
        catatan_menu = page.locator(NIK_R29_SELECTORS["sidebar_blok_catatan"]).first
        catatan_menu.wait_for(state="visible", timeout=20000)
        safe_click(page, catatan_menu, timeout=15000, event_bus=event_bus)
        random_sleep(1.0, 1.5, event_bus=event_bus)
    except Exception as e:
        logger.warning(f"{worker_prefix}Gagal klik menu Blok Catatan: {e}")

    # 2. Pastikan switch toggle anomali aktif di Blok Catatan
    ensure_anomali_switch_active(page, event_bus=event_bus, wait_timeout=35000, worker_prefix=worker_prefix)
    random_sleep(1.5, 2.2, event_bus=event_bus, allow_hesitation=False)

    # 3. Klik Kirim pada formulir Fasih
    logger.bot(f"{worker_prefix}[1/4] Mengklik tombol Kirim formulir Fasih...")
    form_kirim_btn = page.locator(NIK_R29_SELECTORS["btn_kirim"]).first
    form_kirim_btn.wait_for(state="attached", timeout=25000)
    safe_click(page, form_kirim_btn, timeout=25000, event_bus=event_bus, force=True)
    random_sleep(1.0, 1.8, event_bus=event_bus, allow_hesitation=False)

    # 4. Dialog Pop-up Validasi Kirim
    btn_konfirmasi = page.locator(NIK_R29_SELECTORS["btn_konfirmasi_kirim"]).last
    if not btn_konfirmasi.is_visible(timeout=2500):
        logger.bot(f"{worker_prefix}[2/4] Menunggu pop-up validasi Kirim...")
        dialog = page.locator("div[role='dialog'], div[role='alertdialog'], div:has(button:has-text('Batal')):has(button:has-text('Kirim'))").last
        dialog.wait_for(state="visible", timeout=25000)
        random_sleep(0.5, 0.8, event_bus=event_bus, allow_hesitation=False)

        # Cek jumlah Galat
        galat_count = 0
        try:
            galat_card = dialog.locator("button:has-text('Galat')").first
            if galat_card.count() > 0 and galat_card.is_visible():
                card_text = galat_card.inner_text().strip()
                for line in card_text.split("\n"):
                    if line.strip().isdigit():
                        galat_count = int(line.strip())
                        break
                else:
                    nums = re.findall(r"\b\d+\b", card_text)
                    if nums:
                        galat_count = int(nums[0])
            else:
                span_num = dialog.locator("xpath=.//*[text()='Galat']/following-sibling::*[1]").first
                if span_num.count() > 0 and span_num.inner_text().strip().isdigit():
                    galat_count = int(span_num.inner_text().strip())
        except Exception:
            pass

        if galat_count > 0:
            err_msg = f"Terdapat {galat_count} Galat pada form. Submit dibatalkan."
            logger.warning(f"⚠️ {worker_prefix}{err_msg}")
            try:
                btn_batal = dialog.locator("button:has-text('Batal')").first
                if btn_batal.is_visible(timeout=3000):
                    safe_click(page, btn_batal, timeout=5000, event_bus=event_bus)
                else:
                    page.keyboard.press("Escape")
            except Exception:
                page.keyboard.press("Escape")
            return False, f"GAGAL ({galat_count} Galat)", err_msg

        logger.bot(f"{worker_prefix}[2/4] Galat = 0. Mengklik tombol Kirim pada dialog validasi...")
        btn_kirim_dialog = dialog.locator("button:has-text('Kirim')").last
        safe_click(page, btn_kirim_dialog, timeout=20000, event_bus=event_bus)
        random_sleep(1.5, 2.2, event_bus=event_bus, allow_hesitation=False)

    # 5. Dialog Konfirmasi Kirim
    logger.bot(f"{worker_prefix}[3/4] Menunggu dialog Konfirmasi Kirim...")
    btn_konfirmasi = page.locator(NIK_R29_SELECTORS["btn_konfirmasi_kirim"]).last
    btn_konfirmasi.wait_for(state="visible", timeout=30000)
    random_sleep(0.5, 1.0, event_bus=event_bus, allow_hesitation=False)

    if debug_review:
        logger.warning(f"🔍 [MODE REVIEW DEBUGGING] {worker_prefix}Baris {row_idx} [{nama_art}]: Modal Konfirmasi Kirim terbuka.")
        logger.info("Silakan periksa browser. Klik '✅ Konfirmasi Submit & Lanjut' di aplikasi untuk mengirimkan.")
        WindowController.restore_and_focus_browser()
        event_bus.emit("waiting_confirm_submit", {
            "row_idx": row_idx,
            "nama_usaha": nama_art,
            "worker_id": worker_prefix
        })
        event_bus.wait_for_submit_confirmation()
        if event_bus.is_stop_requested():
            return False, "DIHENTIKAN", "Dihentikan oleh pengguna saat review konfirmasi"
        safe_click(page, btn_konfirmasi, timeout=20000, event_bus=event_bus, force=True)
        if auto_minimize:
            WindowController.minimize_browser()
    else:
        logger.bot(f"{worker_prefix}[3/4] Mengklik tombol Konfirmasi...")
        safe_click(page, btn_konfirmasi, timeout=20000, event_bus=event_bus, force=True)

    # 6. Menunggu dialog sukses
    logger.bot(f"{worker_prefix}[4/4] Menunggu respon submit Fasih...")
    try:
        success_dialog = page.locator(NIK_R29_SELECTORS["dialog_sukses"]).first
        success_dialog.wait_for(state="visible", timeout=40000)
        logger.success(f"{worker_prefix}Assignment berhasil disubmit.")
    except Exception:
        logger.warning(f"{worker_prefix}Dialog sukses tidak terdeteksi eksplisit, melanjutkan...")

    return True, "SUKSES", "Assignment berhasil disubmit"

# ==============================================================================
# ENGINE OTOMASI DATA CLEANING NIK R29 (MULTI-WORKER RUNNER)
# ==============================================================================

class SessionExpiredException(Exception):
    """Exception khusus saat sesi Fasih SSO kedaluwarsa atau belum terotentikasi di tengah jalan."""
    pass

class CleaningNIKR29Engine:
    def __init__(self, excel_path, event_bus, config_options=None):
        self.excel_path = excel_path
        self.event_bus = event_bus
        self.config_options = config_options or {}
        self.auto_minimize = self.config_options.get("auto_minimize", True)
        self.debug_review_submit = self.config_options.get("debug_review_submit", False)
        self.row_start = self.config_options.get("row_start", 2)
        self.row_end = self.config_options.get("row_end", None)

        # Konfigurasi Mode Login SSO
        self.auth_mode = self.config_options.get("auth_mode", "manual")  # "manual" | "auto"
        self.sso_username = self.config_options.get("sso_username", "").strip()
        self.sso_password = self.config_options.get("sso_password", "").strip()

        # Konfigurasi Multi-Worker (1 s.d. 5 worker)
        raw_workers = self.config_options.get("num_workers", 2)
        try:
            self.num_workers = max(1, min(5, int(raw_workers)))
        except (ValueError, TypeError):
            self.num_workers = 2

        # Browser & CDP options
        self.browser_type = self.config_options.get("browser_type") or BrowserManager.get_active_browser()
        self.incognito = self.config_options.get("incognito") if self.config_options.get("incognito") is not None else BrowserManager.is_incognito_active()
        self.cdp_url = self.config_options.get("cdp_url") or BrowserManager.get_cdp_url(self.browser_type)

        if self.debug_review_submit and self.num_workers > 1:
            logger.warning("⚠️ Mode Review Debugging aktif: Jumlah worker disetel ke 1 agar review dapat dilakukan teratur.")
            self.num_workers = 1

        # Thread synchronization locks
        self.excel_lock = threading.Lock()
        self.stats_lock = threading.Lock()
        self.reauth_lock = threading.Lock()
        self.task_queue = queue.Queue()

        # Shared session statistics
        self.done_session = 0
        self.error_count = 0
        self.skipped_count = 0
        self.processed_count = 0

    def check_pause_and_stop(self, row_idx=None, audit_data=None, worker_id=None):
        """Memeriksa sinyal jeda atau henti secara responsif di setiap tahapan."""
        if self.event_bus.is_stop_requested():
            if row_idx and audit_data:
                audit_data["status"] = "DIHENTIKAN"
                audit_data["keterangan"] = "Dihentikan oleh pengguna"
                self.event_bus.emit("row_completed", {
                    "row_idx": row_idx,
                    "status": "⏹️ DIHENTIKAN",
                    "keterangan": "Dihentikan oleh pengguna",
                    "worker_id": worker_id
                })
            return False
        self.event_bus.wait_if_paused()
        if self.event_bus.is_stop_requested():
            if row_idx and audit_data:
                audit_data["status"] = "DIHENTIKAN"
                audit_data["keterangan"] = "Dihentikan oleh pengguna"
                self.event_bus.emit("row_completed", {
                    "row_idx": row_idx,
                    "status": "⏹️ DIHENTIKAN",
                    "keterangan": "Dihentikan oleh pengguna",
                    "worker_id": worker_id
                })
            return False
        return True

    def _handle_relogin_needed(self, page, context, row_idx=None, nama_art=None, is_midrun=False, worker_id=1):
        """Menangani kondisi sesi Fasih SSO kedaluwarsa atau belum login."""
        worker_tag = f"[W-{worker_id}] " if worker_id else ""
        with self.reauth_lock:
            if is_midrun:
                try:
                    if page and not page.is_closed() and BrowserManager.is_page_logged_in(page) and not BrowserManager.is_session_expired(page):
                        logger.info(f"{worker_tag}Sesi Fasih SSO telah dipulihkan oleh worker lain. Melanjutkan...")
                        return page
                except Exception:
                    pass

            # Opsi 1: Coba Auto-Login SSO BPS jika mode 'auto' dipilih pengguna
            if self.auth_mode == "auto" and self.sso_username and self.sso_password:
                logger.bot(f"{worker_tag}Sesi belum aktif / kedaluwarsa. Menjalankan Auto-Login SSO BPS (User: {self.sso_username})...")
                login_page = page
                if not login_page or login_page.is_closed():
                    if context.pages:
                        login_page = context.pages[0]
                    else:
                        login_page = context.new_page()

                auto_ok, auto_err = BrowserManager.perform_auto_login(
                    login_page, self.sso_username, self.sso_password
                )
                if auto_ok:
                    logger.success(f"🟢 {worker_tag}Auto-Login SSO BPS berhasil! Sesi telah aktif.")
                    if self.auto_minimize:
                        WindowController.minimize_browser()
                    time.sleep(1.5)
                    return login_page
                else:
                    logger.warning(f"⚠️ {worker_tag}Auto-Login gagal ({auto_err}). Beralih ke mode Login Manual...")

            # Opsi 2 / Fallback: Mode Login Manual (meminta pengguna login di Google Chrome)
            if is_midrun and row_idx:
                prefix = f"Baris {row_idx} [{nama_art or 'ART'}]"
                msg = f"{worker_tag}Sesi Fasih kedaluwarsa di tengah jalan saat memproses {prefix}. Silakan login kembali di Google Chrome, lalu klik '🟢 Saya Sudah Login'."
                logger.warning(f"⚠️ {worker_tag}{prefix}: Sesi Fasih SSO kedaluwarsa di tengah jalan! Menjeda seluruh worker & menampilkan Google Chrome...")
            else:
                prefix = "Pemeriksaan Sesi Awal"
                msg = "Sesi Fasih SSO belum terdeteksi. Silakan login SSO BPS & OTP di Google Chrome, lalu klik '🟢 Saya Sudah Login'."
                logger.warning(f"⚠️ {prefix}: Sesi Fasih SSO belum aktif. Menampilkan jendela Google Chrome...")

            WindowController.restore_and_focus_browser()
            self.event_bus.reset_auth_confirmation()

            self.event_bus.emit("auth_waiting", {"message": msg, "is_midrun": is_midrun, "row_idx": row_idx, "worker_id": worker_id})
            logger.info(f"{worker_tag}Menunggu Anda login SSO BPS di Google Chrome ({'Mid-Run Baris ' + str(row_idx) if is_midrun else 'Awal'})...")

            active_page = page
            while True:
                if self.event_bus.is_stop_requested():
                    logger.warning(f"{worker_tag}Otomasi dihentikan oleh pengguna saat menunggu login SSO.")
                    return None

                if self.event_bus.is_auth_confirmed():
                    logger.success(f"🟢 {worker_tag}Konfirmasi login diterima dari pengguna! Memeriksa tab browser...")
                    detected = BrowserManager.find_authenticated_page(context)
                    if detected:
                        active_page = detected
                    elif not active_page or active_page.is_closed():
                        active_page = context.pages[0] if context.pages else None
                    break

                try:
                    if active_page and not active_page.is_closed():
                        if BrowserManager.is_page_logged_in(active_page) and not BrowserManager.is_session_expired(active_page):
                            logger.success(f"✅ {worker_tag}Login SSO Fasih berhasil dideteksi otomatis pada tab aktif!")
                            break
                except Exception:
                    pass

                detected = BrowserManager.find_authenticated_page(context)
                if detected:
                    active_page = detected
                    logger.success(f"✅ {worker_tag}Login SSO Fasih berhasil dideteksi otomatis pada tab browser!")
                    break

                time.sleep(0.5)

            if self.event_bus.is_stop_requested():
                return None

            self.event_bus.reset_auth_confirmation()

            if not active_page or active_page.is_closed():
                active_page = BrowserManager.find_authenticated_page(context)
                if not active_page and context.pages:
                    active_page = context.pages[0]

            try:
                active_page.wait_for_load_state("domcontentloaded", timeout=6000)
            except Exception:
                pass
            time.sleep(2.0)

            logger.success(f"🟢 {worker_tag}Sesi Fasih aktif! Melanjutkan proses otomasi...")
            self.event_bus.emit("auth_success", {"is_midrun": is_midrun, "row_idx": row_idx, "worker_id": worker_id})

            if self.auto_minimize:
                WindowController.minimize_browser()

            return active_page

    def _navigate_assignment_page(self, page, raw_url, row_idx=None, nama_art="", worker_prefix=""):
        """Membuka URL penugasan dengan proteksi terhadap interupsi navigasi."""
        max_attempts = 4
        last_err = None
        for attempt in range(max_attempts):
            try:
                if page.is_closed():
                    raise Exception("Tab browser tertutup sebelum navigasi URL.")

                try:
                    page.wait_for_load_state("domcontentloaded", timeout=4000)
                except Exception:
                    pass

                page.goto(raw_url, timeout=45000, wait_until="domcontentloaded")
                random_sleep(1.2, 2.0, event_bus=self.event_bus)
                return True
            except Exception as nav_e:
                last_err = nav_e
                nav_err_msg = str(nav_e).lower()
                if any(k in nav_err_msg for k in ["interrupted", "navigation", "abort", "failed to navigate"]):
                    if attempt < max_attempts - 1:
                        logger.warning(
                            f"{worker_prefix}Navigasi baris {row_idx} terinterupsi ({nav_e}). "
                            f"Mencoba ulang ({attempt + 1}/{max_attempts})..."
                        )
                        time.sleep(2.5)
                        continue
                if attempt < max_attempts - 1:
                    time.sleep(1.5)
                else:
                    raise last_err
        return False

    def _process_single_assignment(self, worker_id, page, context, row_idx, wb, sheet, col_map, end_r, target_total):
        """Memproses satu baris penugasan Data Cleaning NIK R29 secara mandiri per worker."""
        w_tag = f"[W-{worker_id}] "

        assignment_col = col_map.get("assignment_id")
        no_urut_col = col_map.get("no_urut_kk")
        nik_col = col_map.get("nomor_induk_kependudukan")
        url_col = col_map.get("url")
        no_kk_col = col_map.get("nomor_kartu_keluarga")
        nama_col = col_map.get("nama")
        status_col = col_map.get("status")
        ket_col = col_map.get("keterangan")

        with self.excel_lock:
            url_val = sheet.cell(row=row_idx, column=url_col).value if url_col else None
            assignment_id = str(sheet.cell(row=row_idx, column=assignment_col).value or "").strip() if assignment_col else ""
            no_urut_kk = str(sheet.cell(row=row_idx, column=no_urut_col).value or "").strip() if no_urut_col else ""
            nik_val = str(sheet.cell(row=row_idx, column=nik_col).value or "").strip() if nik_col else ""
            no_kk_val = str(sheet.cell(row=row_idx, column=no_kk_col).value or "").strip() if no_kk_col else ""
            nama_art = str(sheet.cell(row=row_idx, column=nama_col).value or "").strip() if nama_col else ""
            status_val = str(sheet.cell(row=row_idx, column=status_col).value or "").strip() if status_col else ""

        if not url_val or str(url_val).strip() == "":
            return

        raw_url = str(url_val).strip()
        status_upper = status_val.strip().upper()

        # Lewati jika sudah selesai pada sesi sebelumnya
        if "SUKSES" in status_upper or "DONE" in status_upper:
            with self.stats_lock:
                self.skipped_count += 1
                self.processed_count += 1
                current_proc = self.processed_count
            logger.info(f"⏭️ {w_tag}Baris {row_idx} [{nama_art or 'ART'}] sudah berstatus '{status_val}'. Dilewati.")
            self.event_bus.emit("progress_updated", {
                "current": current_proc,
                "total": target_total,
                "row_idx": row_idx
            })
            return

        # Setup audit data record
        audit_data = {
            "row_idx": row_idx,
            "nama_art": nama_art,
            "assignment_id": assignment_id,
            "no_urut_kk": no_urut_kk,
            "nik": nik_val,
            "no_kk": no_kk_val,
            "url": raw_url,
            "status": "PROSES",
            "keterangan": "",
            "timestamp": datetime.now().isoformat(),
            "worker_id": worker_id
        }

        self.event_bus.emit("worker_status_updated", {
            "worker_id": worker_id,
            "status": "BUSY",
            "row_idx": row_idx,
            "nama_usaha": f"{nama_art[:18]} (NIK: {nik_val[:10]}...)"
        })

        self.event_bus.emit("row_started", {
            "row_idx": row_idx,
            "nama_usaha": nama_art,
            "url": raw_url,
            "worker_id": worker_id
        })

        logger.bot(f"{w_tag}--- [Baris {row_idx}/{end_r}] Memproses NIK: {nama_art} (Urutan: {no_urut_kk}, NIK: {nik_val}) ---")

        reauth_count = 0
        max_reauth_attempts = 2

        try:
            while reauth_count <= max_reauth_attempts:
                if not self.check_pause_and_stop(row_idx, audit_data, worker_id=worker_id):
                    return

                try:
                    # STEP 1: Navigasi ke URL penugasan
                    self._navigate_assignment_page(page, raw_url, row_idx=row_idx, nama_art=nama_art, worker_prefix=w_tag)

                    # Pengecekan sesi setelah halaman penugasan dimuat
                    if BrowserManager.is_session_expired(page) or not BrowserManager.is_page_logged_in(page):
                        raise SessionExpiredException(f"{w_tag}Sesi Fasih kedaluwarsa saat membuka form Baris {row_idx}")

                    # STEP 2 & 3 & 4 & 5: Menuju Blok I & Cari Anggota Keluarga yang match Nama & No Urut
                    if not self.check_pause_and_stop(row_idx, audit_data, worker_id=worker_id):
                        return

                    art_ok, matched_art_name, err_art = _find_and_open_matching_art(
                        page,
                        excel_no_urut=no_urut_kk,
                        excel_nama=nama_art,
                        event_bus=self.event_bus,
                        worker_prefix=w_tag,
                        check_stop_fn=lambda: not self.event_bus.is_stop_requested()
                    )

                    if not art_ok:
                        # Lewati / gagal jika tidak match
                        logger.warning(f"⚠️ {w_tag}Baris {row_idx} [{nama_art}]: {err_art}. Melewati assignment.")
                        with self.excel_lock:
                            sheet.cell(row=row_idx, column=status_col).value = "DILEWATI"
                            sheet.cell(row=row_idx, column=ket_col).value = err_art[:150]
                            try:
                                wb.save(self.excel_path)
                            except Exception:
                                pass

                        with self.stats_lock:
                            self.skipped_count += 1
                            self.processed_count += 1
                            current_proc = self.processed_count

                        audit_data["status"] = "DILEWATI"
                        audit_data["keterangan"] = err_art
                        audit_logger.record_row(audit_data)

                        self.event_bus.emit("row_completed", {
                            "row_idx": row_idx,
                            "status": "⏭️ DILEWATI",
                            "keterangan": err_art[:100],
                            "worker_id": worker_id
                        })
                        self.event_bus.emit("progress_updated", {
                            "current": current_proc,
                            "total": target_total,
                            "row_idx": row_idx
                        })
                        return

                    # STEP 6 & 7 & 8 & 9: Cek rincian NIK (nik_dtsen) vs Excel & Berikan Remarks '#3' jika sama
                    if not self.check_pause_and_stop(row_idx, audit_data, worker_id=worker_id):
                        return

                    nik_cocok, web_nik_val, err_nik = _verify_and_remark_nik(
                        page,
                        excel_nik=nik_val,
                        event_bus=self.event_bus,
                        worker_prefix=w_tag
                    )

                    if not nik_cocok:
                        # Jika NIK Beda -> Skip ke assignment selanjutnya dengan keterangan
                        logger.warning(f"⚠️ {w_tag}Baris {row_idx} [{nama_art}]: {err_nik}. Melewati assignment.")
                        with self.excel_lock:
                            sheet.cell(row=row_idx, column=status_col).value = "DILEWATI"
                            sheet.cell(row=row_idx, column=ket_col).value = err_nik[:150]
                            try:
                                wb.save(self.excel_path)
                            except Exception:
                                pass

                        with self.stats_lock:
                            self.skipped_count += 1
                            self.processed_count += 1
                            current_proc = self.processed_count

                        audit_data["status"] = "DILEWATI"
                        audit_data["keterangan"] = err_nik
                        audit_logger.record_row(audit_data)

                        self.event_bus.emit("row_completed", {
                            "row_idx": row_idx,
                            "status": "⏭️ DILEWATI",
                            "keterangan": err_nik[:100],
                            "worker_id": worker_id
                        })
                        self.event_bus.emit("progress_updated", {
                            "current": current_proc,
                            "total": target_total,
                            "row_idx": row_idx
                        })
                        return

                    # STEP 10: Lanjut pengecekan aktif tampilkan anomali di Blok Catatan sampai akhir submit
                    if not self.check_pause_and_stop(row_idx, audit_data, worker_id=worker_id):
                        return

                    sub_ok, final_status, final_ket = _submit_assignment_form(
                        page,
                        self.event_bus,
                        row_idx=row_idx,
                        nama_art=nama_art,
                        debug_review=self.debug_review_submit,
                        auto_minimize=self.auto_minimize,
                        worker_prefix=w_tag
                    )

                    # Simpan hasil akhir ke Excel secara thread-safe
                    with self.excel_lock:
                        sheet.cell(row=row_idx, column=status_col).value = final_status
                        sheet.cell(row=row_idx, column=ket_col).value = final_ket[:150]
                        try:
                            wb.save(self.excel_path)
                        except Exception as save_err:
                            logger.warning(f"⚠️ {w_tag}Gagal menyimpan checkpoint Excel baris {row_idx}: {save_err}")

                    with self.stats_lock:
                        if sub_ok:
                            self.done_session += 1
                        else:
                            self.error_count += 1
                        self.processed_count += 1
                        current_proc = self.processed_count

                    audit_data["status"] = final_status
                    audit_data["keterangan"] = final_ket
                    audit_logger.record_row(audit_data)

                    badge_status = f"✅ {final_status}" if sub_ok else f"❌ {final_status}"
                    self.event_bus.emit("row_completed", {
                        "row_idx": row_idx,
                        "status": badge_status,
                        "keterangan": final_ket[:100],
                        "worker_id": worker_id
                    })
                    self.event_bus.emit("progress_updated", {
                        "current": current_proc,
                        "total": target_total,
                        "row_idx": row_idx
                    })
                    break

                except SessionExpiredException as sess_err:
                    reauth_count += 1
                    logger.warning(f"⚠️ {w_tag}Baris {row_idx}: {sess_err}. Menjeda worker & memicu re-login...")
                    self.event_bus.request_pause()
                    page = self._handle_relogin_needed(page, context, row_idx=row_idx, nama_art=nama_art, is_midrun=True, worker_id=worker_id)
                    self.event_bus.request_resume()

                    if not page or self.event_bus.is_stop_requested():
                        audit_data["status"] = "DIHENTIKAN"
                        audit_data["keterangan"] = "Dihentikan oleh pengguna saat sesi login kedaluwarsa"
                        self.event_bus.emit("row_completed", {
                            "row_idx": row_idx,
                            "status": "⏹️ DIHENTIKAN",
                            "keterangan": "Dihentikan oleh pengguna saat sesi login kedaluwarsa",
                            "worker_id": worker_id
                        })
                        return

                    logger.info(f"🔄 {w_tag}Mengulang kembali proses otomasi Baris {row_idx} [{nama_art[:20]}] dari awal...")
                    time.sleep(2.0)
                    continue

                except Exception as step_err:
                    if (BrowserManager.is_session_expired(page) or not BrowserManager.is_page_logged_in(page)) and reauth_count < max_reauth_attempts:
                        reauth_count += 1
                        logger.warning(f"⚠️ {w_tag}Baris {row_idx}: Sesi terputus ({step_err}). Menjeda worker & memicu re-login...")
                        self.event_bus.request_pause()
                        page = self._handle_relogin_needed(page, context, row_idx=row_idx, nama_art=nama_art, is_midrun=True, worker_id=worker_id)
                        self.event_bus.request_resume()

                        if not page or self.event_bus.is_stop_requested():
                            audit_data["status"] = "DIHENTIKAN"
                            audit_data["keterangan"] = "Dihentikan oleh pengguna saat sesi login kedaluwarsa"
                            self.event_bus.emit("row_completed", {
                                "row_idx": row_idx,
                                "status": "⏹️ DIHENTIKAN",
                                "keterangan": "Dihentikan oleh pengguna saat sesi login kedaluwarsa",
                                "worker_id": worker_id
                            })
                            return

                        logger.info(f"🔄 {w_tag}Mengulang kembali proses otomasi Baris {row_idx} [{nama_art[:20]}] dari awal...")
                        time.sleep(2.0)
                        continue
                    else:
                        raise step_err

        except Exception as row_err:
            error_msg = str(row_err)
            logger.error(f"❌ {w_tag}Baris {row_idx} [{nama_art[:20]}] GAGAL: {error_msg}")
            with self.excel_lock:
                sheet.cell(row=row_idx, column=status_col).value = "GAGAL"
                sheet.cell(row=row_idx, column=ket_col).value = error_msg[:150]
                try:
                    wb.save(self.excel_path)
                except Exception:
                    pass

            with self.stats_lock:
                self.error_count += 1

            audit_data["status"] = "GAGAL"
            audit_data["keterangan"] = error_msg
            audit_logger.record_row(audit_data)

            self.event_bus.emit("row_completed", {
                "row_idx": row_idx,
                "status": "❌ GAGAL",
                "keterangan": error_msg[:100],
                "worker_id": worker_id
            })

        finally:
            self.event_bus.emit("worker_status_updated", {
                "worker_id": worker_id,
                "status": "IDLE",
                "row_idx": None,
                "nama_usaha": ""
            })
            random_sleep(0.8, 1.5, event_bus=self.event_bus)

    def run(self):
        """Metode utama eksekusi engine Data Cleaning NIK R29."""
        logger.info(f"Memulai Engine Data Cleaning NIK R29 SE2026 (Konfigurasi: {self.num_workers} Worker)...")
        self.event_bus.set_state(EngineState.RUNNING)

        wb = None

        # 1. Validasi awal template Excel
        val_res = ExcelNIKR29Validator.validate_excel_file(self.excel_path)
        if not val_res["valid"]:
            logger.error(f"Validasi file Excel gagal: {val_res['error']}")
            self.event_bus.set_state(EngineState.IDLE)
            self.event_bus.emit("all_completed", {
                "done": 0, "error": 0, "skipped": 0, "was_stopped": False
            })
            return

        col_map = val_res["col_map"]

        try:
            wb = openpyxl.load_workbook(self.excel_path)
            sheet = wb.active
        except Exception as e:
            logger.error(f"Gagal memuat workbook Excel: {e}")
            self.event_bus.set_state(EngineState.IDLE)
            self.event_bus.emit("all_completed", {
                "done": 0, "error": 0, "skipped": 0, "was_stopped": False
            })
            return

        # 2. Hubungkan ke Browser CDP (atau luncurkan otomatis jika belum aktif)
        cdp_port = BrowserManager.get_cdp_port(self.browser_type)
        b_label = "Microsoft Edge" if self.browser_type == "edge" else "Google Chrome"
        is_running, b_name = BrowserManager.is_cdp_running(browser_type=self.browser_type, port=cdp_port)
        if not is_running:
            logger.info(f"Browser ({b_label}) belum aktif di port {cdp_port}. Meluncurkan {b_label} secara otomatis...")
            launched, launch_msg = BrowserManager.launch_browser(browser_type=self.browser_type, incognito=self.incognito)
            if not launched:
                logger.error(f"Gagal meluncurkan browser: {launch_msg}")
                self.event_bus.set_state(EngineState.IDLE)
                self.event_bus.emit("all_completed", {
                    "done": 0, "error": 0, "skipped": 0, "was_stopped": False
                })
                wb.close()
                return
            time.sleep(1.5)

        # 3. Probe SSO Awal
        logger.info("Memeriksa status otentikasi SSO Fasih awal...")
        try:
            with sync_playwright() as p_probe:
                browser_probe = p_probe.chromium.connect_over_cdp(self.cdp_url)
                context_probe = browser_probe.contexts[0]
                page_probe = BrowserManager.find_authenticated_page(context_probe)
                if not page_probe:
                    for p_cand in context_probe.pages:
                        if "fasih-sm.bps.go.id" in p_cand.url:
                            page_probe = p_cand
                            break
                    if not page_probe:
                        page_probe = context_probe.pages[0] if context_probe.pages else context_probe.new_page()
                        if "fasih-sm.bps.go.id" not in page_probe.url:
                            try:
                                page_probe.goto("https://fasih-sm.bps.go.id/app", timeout=30000)
                            except Exception:
                                pass

                if not BrowserManager.is_page_logged_in(page_probe) or BrowserManager.is_session_expired(page_probe):
                    page_probe = self._handle_relogin_needed(page_probe, context_probe, is_midrun=False, worker_id=1)
                    if not page_probe or self.event_bus.is_stop_requested():
                        logger.warning("Otomasi dibatalkan sebelum login SSO selesai.")
                        wb.close()
                        self.event_bus.set_state(EngineState.IDLE)
                        self.event_bus.emit("all_completed", {
                            "done": 0, "error": 0, "skipped": 0, "was_stopped": True
                        })
                        return
                elif self.auto_minimize:
                    WindowController.minimize_browser()
        except Exception as e_probe:
            logger.error(f"Gagal memeriksa status SSO Fasih: {e_probe}")
            wb.close()
            self.event_bus.set_state(EngineState.IDLE)
            self.event_bus.emit("all_completed", {
                "done": 0, "error": 0, "skipped": 0, "was_stopped": False
            })
            return

        # 4. Kolom mapping & persiapan checkpoint
        url_col = col_map.get("url")
        status_col = col_map.get("status")
        ket_col = col_map.get("keterangan")

        with self.excel_lock:
            if not status_col:
                status_col = sheet.max_column + 1
                col_map["status"] = status_col
                sheet.cell(row=1, column=status_col).value = "Status_Eksekusi"
                logger.info(f"Kolom 'Status_Eksekusi' otomatis dibuat di kolom {status_col}.")
            if not ket_col:
                ket_col = sheet.max_column + 1
                col_map["keterangan"] = ket_col
                sheet.cell(row=1, column=ket_col).value = "Keterangan"
                logger.info(f"Kolom 'Keterangan' otomatis dibuat di kolom {ket_col}.")

        max_row = sheet.max_row
        start_r = max(2, self.row_start)
        end_r = min(max_row, self.row_end) if self.row_end else max_row

        target_rows = [
            r for r in range(start_r, end_r + 1)
            if url_col and sheet.cell(row=r, column=url_col).value and str(sheet.cell(row=r, column=url_col).value).strip() != ""
        ]
        target_total = len(target_rows) if target_rows else max(1, end_r - start_r + 1)

        if not target_rows:
            logger.warning("Tidak ditemukan baris penugasan yang valid dalam rentang yang dipilih.")
            wb.close()
            self.event_bus.set_state(EngineState.IDLE)
            self.event_bus.emit("all_completed", {
                "done": 0, "error": 0, "skipped": 0, "was_stopped": False
            })
            return

        actual_workers = min(self.num_workers, len(target_rows))
        logger.info(f"Rentang baris target: {start_r} s.d. {end_r} (Ditemukan {len(target_rows)} baris penugasan). Mengaktifkan {actual_workers} Worker.")

        for r_idx in target_rows:
            self.task_queue.put(r_idx)

        # Worker runner loop
        def worker_runner(w_id):
            w_tag = f"[W-{w_id}] "
            try:
                with sync_playwright() as p_worker:
                    browser = p_worker.chromium.connect_over_cdp(self.cdp_url)
                    context = browser.contexts[0]

                    # Buka tab khusus worker
                    worker_page = context.new_page()
                    try:
                        worker_page.wait_for_load_state("domcontentloaded", timeout=4000)
                    except Exception:
                        pass

                    logger.info(f"{w_tag}Tab worker NIK siap beroperasi.")

                    while not self.task_queue.empty():
                        if self.event_bus.is_stop_requested():
                            break

                        try:
                            row_idx = self.task_queue.get_nowait()
                        except queue.Empty:
                            break

                        try:
                            self._process_single_assignment(
                                w_id, worker_page, context,
                                row_idx, wb, sheet, col_map, end_r, target_total
                            )
                        finally:
                            self.task_queue.task_done()

                    try:
                        if not worker_page.is_closed():
                            worker_page.close()
                    except Exception:
                        pass
            except Exception as w_err:
                logger.error(f"❌ {w_tag}Worker mengalami galat kritis: {w_err}")

        # Jalankan thread worker
        worker_threads = []
        for w_i in range(1, actual_workers + 1):
            t = threading.Thread(
                target=worker_runner,
                args=(w_i,),
                daemon=True,
                name=f"CleaningNIKR29-Worker-{w_i}"
            )
            worker_threads.append(t)
            t.start()
            time.sleep(0.4)

        for t in worker_threads:
            t.join()

        try:
            wb.close()
        except Exception:
            pass

        was_stopped = self.event_bus.is_stop_requested()
        self.event_bus.set_state(EngineState.IDLE)
        logger.info(f"Otomasi Data Cleaning NIK R29 SE2026 selesai. Berhasil: {self.done_session}, Gagal: {self.error_count}, Dilewati: {self.skipped_count}")

        self.event_bus.emit("all_completed", {
            "done": self.done_session,
            "error": self.error_count,
            "skipped": self.skipped_count,
            "was_stopped": was_stopped
        })

CleaningNIKR29Engine = CleaningNIKR29Engine


CleaningNIKEngine = CleaningNIKR29Engine
