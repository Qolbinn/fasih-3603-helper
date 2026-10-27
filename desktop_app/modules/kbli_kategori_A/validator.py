import os
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from modules.kbli_kategori_A.selectors import EXCEL_KBLI_KATEGORI_A_COLUMNS_CONFIG

class ExcelKBLIKategoriAValidator:
    MANDATORY_KEYS = [
        "keg_utama",
        "kbli_akhir",
        "url",
        "status_kbli",
        "kbli_terbaru"
    ]

    @staticmethod
    def is_file_locked(file_path):
        """Memeriksa apakah file Excel sedang terbuka di aplikasi lain (Microsoft Excel) yang mengunci penulisan."""
        if not os.path.exists(file_path):
            return False
        try:
            with open(file_path, "r+b"):
                return False
        except (IOError, PermissionError):
            return True

    @classmethod
    def validate_excel_file(cls, file_path):
        """
        Melakukan validasi menyeluruh terhadap file Excel input KBLI Kategori A.
        Mengembalikan dict hasil pemeriksaan & data pratinjau.
        """
        p = Path(file_path)
        if not p.exists():
            return {
                "valid": False,
                "error": f"File tidak ditemukan: {file_path}",
                "warning": None,
                "file_locked": False,
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        locked = cls.is_file_locked(file_path)
        warning_msg = None
        if locked:
            warning_msg = "Perhatian: File sedang dibuka di Excel. Anda dapat melihat pratinjau, namun harap tutup file Excel sebelum memulai otomasi agar hasil dapat disimpan."

        try:
            with open(file_path, "rb") as f:
                wb = openpyxl.load_workbook(f, data_only=True)
                sheet = wb.active
        except Exception as e:
            return {
                "valid": False,
                "error": f"Gagal membaca format file Excel: {e}",
                "warning": None,
                "file_locked": locked,
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        # 1. Baca header baris 1
        header_map = {}
        for col_idx in range(1, sheet.max_column + 1):
            val = sheet.cell(row=1, column=col_idx).value
            if val is not None:
                header_map[str(val).strip().lower()] = col_idx

        # 2. Mapping konfigurasi ke index kolom sebenarnya
        col_map = {}
        for key, aliases in EXCEL_KBLI_KATEGORI_A_COLUMNS_CONFIG.items():
            for alias in aliases:
                clean_alias = str(alias).strip().lower()
                if clean_alias in header_map:
                    col_map[key] = header_map[clean_alias]
                    break

        # 3. Periksa kolom wajib
        missing_mandatory = []
        for mand in cls.MANDATORY_KEYS:
            if mand not in col_map:
                missing_mandatory.append(mand)

        if missing_mandatory:
            wb.close()
            return {
                "valid": False,
                "error": f"Template tidak sesuai. Kolom wajib berikut tidak ditemukan: {', '.join(missing_mandatory)}",
                "warning": warning_msg,
                "file_locked": locked,
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        # 4. Validasi isi baris data (mulai baris 2)
        total_rows = 0
        done_rows = 0
        pending_rows = 0
        rows_data = []

        status_col = col_map.get("status")
        ket_col = col_map.get("keterangan")

        for r_idx in range(2, sheet.max_row + 1):
            keg_val = sheet.cell(row=r_idx, column=col_map["keg_utama"]).value
            kbli_lama_val = sheet.cell(row=r_idx, column=col_map["kbli_akhir"]).value
            url_val = sheet.cell(row=r_idx, column=col_map["url"]).value
            st_kbli_val = sheet.cell(row=r_idx, column=col_map["status_kbli"]).value
            kbli_baru_val = sheet.cell(row=r_idx, column=col_map["kbli_terbaru"]).value

            # Jika baris benar-benar kosong, lewati
            if all(v is None or str(v).strip() == "" for v in [keg_val, kbli_lama_val, url_val, st_kbli_val, kbli_baru_val]):
                continue

            total_rows += 1

            # Format nilai pembersihan
            keg_str = str(keg_val).strip() if keg_val is not None else ""
            kbli_lama_str = str(kbli_lama_val).strip() if kbli_lama_val is not None else ""
            if kbli_lama_str.endswith(".0"):
                kbli_lama_str = kbli_lama_str[:-2]

            url_str = str(url_val).strip() if url_val is not None else ""
            st_kbli_str = str(st_kbli_val).strip() if st_kbli_val is not None else ""
            kbli_baru_str = str(kbli_baru_val).strip() if kbli_baru_val is not None else ""
            if kbli_baru_str.endswith(".0"):
                kbli_baru_str = kbli_baru_str[:-2]

            # Cek status & keterangan eksekusi sebelumnya
            status_exec = ""
            if status_col:
                s_val = sheet.cell(row=r_idx, column=status_col).value
                status_exec = str(s_val).strip() if s_val is not None else ""

            ket_exec = ""
            if ket_col:
                k_val = sheet.cell(row=r_idx, column=ket_col).value
                ket_exec = str(k_val).strip() if k_val is not None else ""

            is_done = status_exec in ("SUKSES", "DILEWATI") or "berhasil" in status_exec.lower()
            if is_done:
                done_rows += 1
            else:
                pending_rows += 1

            row_status_display = "Sudah Selesai" if is_done else "Belum Dikerjakan"

            rows_data.append({
                "row_idx": r_idx,
                "keg_utama": keg_str,
                "kbli_akhir": kbli_lama_str,
                "url": url_str,
                "status_kbli": st_kbli_str,
                "kbli_terbaru": kbli_baru_str,
                "status": status_exec,
                "status_eksekusi": status_exec,
                "keterangan": ket_exec,
                "is_done": is_done,
                "status_display": row_status_display
            })

        wb.close()

        if total_rows == 0:
            return {
                "valid": False,
                "error": "File Excel tidak memiliki baris data (hanya header).",
                "warning": warning_msg,
                "file_locked": locked,
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        return {
            "valid": True,
            "error": None,
            "warning": warning_msg,
            "file_locked": locked,
            "total_rows": total_rows,
            "done_rows": done_rows,
            "pending_rows": pending_rows,
            "rows_data": rows_data,
            "col_map": col_map
        }

    @classmethod
    def generate_template(cls, target_path):
        """Membuat file template Excel standar untuk KBLI Kategori A."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data_KBLI_Kategori_A"

        headers = [
            ("keg_utama", "Kegiatan Utama"),
            ("kbli_akhir", "KBLI Akhir"),
            ("url", "Link Penugasan Fasih (URL)"),
            ("Status_KBLI", "Status KBLI"),
            ("KBLI_Terbaru", "KBLI Terbaru"),
            ("Status_Eksekusi", "Status Eksekusi (Opsional)"),
            ("Keterangan", "Keterangan (Opsional)")
        ]

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="047857", end_color="047857", fill_type="solid")  # Emerald/Forest Green Theme (KBLI Kategori A)
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        thin_border = Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="thin", color="CCCCCC")
        )

        ws.row_dimensions[1].height = 28

        for col_idx, (col_id, col_label) in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_id)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        # Contoh baris data sampel KBLI Kategori A (Kombinasi Non-Jasa & Jasa Pertanian)
        sample_rows = [
            (
                "Pertanian Padi Hibrida",
                "01111",
                "https://fasih-sm.bps.go.id/assignment/detail/sample-a-1",
                "Tetap",
                "",
                "",
                ""
            ),
            (
                "JASA PENGGILINGAN PADI",
                "01611",
                "https://fasih-sm.bps.go.id/assignment/detail/sample-a-2",
                "Tetap",
                "",
                "",
                ""
            ),
            (
                "Budidaya Ikan Lele di Kolam Terpal",
                "03221",
                "https://fasih-sm.bps.go.id/assignment/detail/sample-a-3",
                "Ganti",
                "03222",
                "",
                ""
            )
        ]

        data_font = Font(name="Calibri", size=10)
        data_align = Alignment(vertical="center")

        for r_idx, row_values in enumerate(sample_rows, 2):
            ws.row_dimensions[r_idx].height = 22
            for c_idx, val in enumerate(row_values, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = data_font
                cell.alignment = data_align
                cell.border = thin_border

        # Atur lebar kolom otomatis
        col_widths = {
            1: 45,  # keg_utama
            2: 14,  # kbli_akhir
            3: 45,  # url
            4: 16,  # Status_KBLI
            5: 16,  # KBLI_Terbaru
            6: 18,  # Status_Eksekusi
            7: 25   # Keterangan
        }
        for col_idx, width in col_widths.items():
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

        # Data validation dropdown untuk kolom Status_KBLI (D2:D1000)
        dv = DataValidation(type="list", formula1='"Tetap,Ganti"', allow_blank=True)
        dv.error = 'Status_KBLI hanya boleh diisi "Tetap" atau "Ganti"'
        dv.errorTitle = 'Nilai Tidak Valid'
        dv.prompt = 'Pilih "Tetap" atau "Ganti"'
        dv.promptTitle = 'Status KBLI'
        ws.add_data_validation(dv)
        dv.add("D2:D1000")

        ws.freeze_panes = "A2"
        wb.save(target_path)
        wb.close()


# Alias untuk kompatibilitas
ExcelKBLIANonJasaValidator = ExcelKBLIKategoriAValidator
