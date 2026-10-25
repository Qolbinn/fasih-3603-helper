import os
import re
import time
import random
import threading
import queue
import openpyxl
from datetime import datetime
from pathlib import Path
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

from config import CDP_URL
from core.logger import logger
from core.event_bus import EngineState
from core.audit_logger import audit_logger
from core.window_controller import WindowController
from core.browser_manager import BrowserManager
from modules.cleaning_usaha.selectors import (
    PENGELUARAN_MAP,
    PENDAPATAN_MAP,
    ASET_MAP,
    TOTAL_PENGELUARAN_SELECTOR,
    TOTAL_PENDAPATAN_SELECTOR,
    BLOK_II_XPATH,
    BLOK_CATATAN_XPATH,
    KIRIM_BUTTON_XPATH,
    CONFIRM_BUTTON_XPATH,
    RINGKASAN_KIRIM_SELECTOR,
    KONFIRMASI_KIRIM_BUTTON_SELECTOR,
    SUCCESS_SUBMIT_DIALOG_SELECTOR
)
from modules.cleaning_usaha.validator import ExcelValidator

# ==============================================================================
# HELPER BEHAVIOR MANUSIA & PLAYWRIGHT
# ==============================================================================

def responsive_sleep(min_sec=1.0, max_sec=2.0, event_bus=None, allow_hesitation=True):
    duration = random.uniform(min_sec, max_sec)
    if allow_hesitation and random.random() < 0.2:
        duration += random.uniform(1.0, 2.0)
    end_time = time.time() + duration
    while time.time() < end_time:
        if event_bus and event_bus.is_stop_requested():
            return False
        if event_bus and event_bus.is_paused():
            event_bus.wait_if_paused()
            if event_bus.is_stop_requested():
                return False
        time.sleep(min(0.08, max(0.01, end_time - time.time())))
    return True

def human_jitter(min_sec=0.1, max_sec=0.25, event_bus=None):
    responsive_sleep(min_sec, max_sec, event_bus=event_bus, allow_hesitation=False)

def random_sleep(min_sec=1.0, max_sec=2.0, event_bus=None, allow_hesitation=True):
    responsive_sleep(min_sec, max_sec, event_bus=event_bus, allow_hesitation=allow_hesitation)

def safe_click(page, selector_or_locator, timeout=20000, event_bus=None, force=False):
    human_jitter(0.1, 0.2, event_bus=event_bus)
    if isinstance(selector_or_locator, str):
        loc = page.locator(selector_or_locator).first
    else:
        loc = selector_or_locator
    loc.wait_for(state="attached", timeout=timeout)
    try:
        loc.scroll_into_view_if_needed(timeout=5000)
    except Exception:
        pass
    human_jitter(0.1, 0.2, event_bus=event_bus)
    try:
        loc.click(timeout=timeout, force=force)
    except Exception as e:
        err_msg = str(e).lower()
        if "intercepts pointer events" in err_msg or "subtree intercepts" in err_msg or not force:
            logger.bot("Intersepsi klik terdeteksi (animasi/backdrop modal). Mengulang klik dengan force=True...")
            human_jitter(0.1, 0.2, event_bus=event_bus)
            loc.click(timeout=8000, force=True)
        else:
            raise e

def safe_fill_currency(page, container_locator, value, event_bus=None):
    val_int = int(round(float(value)))
    val_str = str(val_int)
    try:
        container_locator.scroll_into_view_if_needed(timeout=5000)
    except Exception:
        pass
    input_elem = container_locator.locator("input[type='text'], input").first
    input_elem.wait_for(state="visible", timeout=25000)
    input_elem.click()
    human_jitter(0.05, 0.15, event_bus=event_bus)
    # Direct fill Playwright: mengosongkan dan mengisi teks langsung ke elemen DOM
    input_elem.fill(val_str)
    human_jitter(0.05, 0.15, event_bus=event_bus)
    try:
        input_elem.dispatch_event("input")
        input_elem.dispatch_event("change")
    except Exception:
        pass
    try:
        page.keyboard.press("Tab")
    except Exception:
        pass

def add_remarks_to_element(page, container_selector, note_text, event_bus=None, worker_prefix=""):
    """Membuka modal Catatan pada elemen container, mengetik catatan, klik simpan, lalu tutup modal."""
    container = page.locator(container_selector).first
    container.wait_for(state="attached", timeout=30000)
    try:
        container.scroll_into_view_if_needed(timeout=5000)
    except Exception:
        pass
    random_sleep(0.4, 0.7, event_bus=event_bus)

    # Cari tombol Catatan yang VISIBLE di dalam container
    btn_catatan = container.locator("button[title='Catatan']:visible, button[aria-label='Catatan']:visible").first
    if btn_catatan.count() == 0 or not btn_catatan.is_visible():
        all_btns = container.locator("button[title='Catatan'], button[aria-label='Catatan']")
        btn_catatan = None
        for i in range(all_btns.count()):
            b = all_btns.nth(i)
            if b.is_visible():
                btn_catatan = b
                break

    if not btn_catatan:
        raise Exception(f"Tombol Catatan visible tidak ditemukan pada container {container_selector}")

    safe_click(page, btn_catatan, timeout=20000, event_bus=event_bus)
    random_sleep(0.7, 1.1, event_bus=event_bus)

    dialog = page.locator("div[role='dialog']").first
    try:
        dialog.wait_for(state="visible", timeout=25000)
    except PlaywrightTimeoutError:
        raise Exception("Modal pop-up Catatan tidak muncul setelah tombol Catatan diklik")

    note_str = str(note_text).strip()

    # Cek apakah catatan sudah pernah tercatat sebelumnya di riwayat pesan
    already_exists = False
    try:
        bubbles = dialog.locator("div[class*='whitespace-pre-wrap'], div[class*='wrap-break-word']")
        for i in range(bubbles.count()):
            if bubbles.nth(i).inner_text().strip() == note_str:
                already_exists = True
                break
    except Exception:
        pass

    if already_exists:
        logger.info(f"{worker_prefix}Catatan '{note_str}' sudah tercatat sebelumnya. Melewati pengiriman ulang.")
    else:
        textarea = dialog.locator("textarea[placeholder*='Tambah catatan'], textarea").first
        try:
            textarea.wait_for(state="visible", timeout=25000)
        except PlaywrightTimeoutError:
            raise Exception("Textarea Catatan tidak ditemukan di dalam dialog")

        textarea.click()
        human_jitter(0.1, 0.2, event_bus=event_bus)
        textarea.fill(note_str)
        human_jitter(0.15, 0.25, event_bus=event_bus)

        btn_simpan = dialog.locator("button[title='Simpan']").first
        btn_simpan.wait_for(state="visible", timeout=25000)
        safe_click(page, btn_simpan, timeout=20000, event_bus=event_bus)
        random_sleep(0.8, 1.2, event_bus=event_bus)

    try:
        btn_close = dialog.locator("button[aria-label='Dismiss'], button[title='Tutup'], button:has-text('Tutup')").first
        if btn_close.is_visible(timeout=4000):
            safe_click(page, btn_close, timeout=8000, event_bus=event_bus)
        else:
            page.keyboard.press("Escape")
    except Exception:
        page.keyboard.press("Escape")

    try:
        dialog.wait_for(state="hidden", timeout=8000)
    except Exception:
        page.keyboard.press("Escape")

    random_sleep(0.4, 0.7, event_bus=event_bus)

def ensure_anomali_switch_active(page, event_bus=None, wait_timeout=30000, worker_prefix=""):
    """
    Pastikan switch 'Tampilkan Anomali Usaha dan Keluarga' (div#cek_anomali_button) AKTIF.
    - Sudah aktif (input aria-checked='true')  -> dibiarkan, tidak diklik.
    - Belum aktif (aria-checked='false')       -> klik kontrol switch, lalu verifikasi benar-benar aktif.
    Raise Exception jika elemen tidak muncul atau gagal diaktifkan.
    """
    box = page.locator("div#cek_anomali_button").first
    try:
        # Blok Catatan baru dibuka -> tunggu elemen benar-benar dirender (bukan sekadar count() == 0 lalu dilewati)
        box.wait_for(state="visible", timeout=wait_timeout)
    except PlaywrightTimeoutError:
        raise Exception("Switch 'Tampilkan Anomali Usaha dan Keluarga' tidak muncul di Blok Catatan.")

    try:
        box.scroll_into_view_if_needed(timeout=5000)
    except Exception:
        pass
    sw_input = box.locator("input[role='switch']").first
    sw_control = box.locator("[id^='switch-'][id$='-control']").first

    def _is_active():
        return sw_input.get_attribute("aria-checked") == "true"

    if _is_active():
        logger.bot(f"{worker_prefix}Switch 'Tampilkan Anomali Usaha dan Keluarga' sudah aktif -> Dibiarkan.")
        return False

    logger.bot(f"{worker_prefix}Switch 'Tampilkan Anomali Usaha dan Keluarga' belum aktif -> Mengaktifkan...")
    for attempt in range(1, 4):
        # Input asli disembunyikan (1px, clip) -> yang diklik adalah div kontrol switch-nya
        safe_click(page, sw_control, timeout=5000, event_bus=event_bus)

        # Tunggu state berubah (maks ~3 detik)
        for _ in range(10):
            random_sleep(0.25, 0.35, event_bus=event_bus, allow_hesitation=False)
            if _is_active():
                logger.bot(f"{worker_prefix}Switch 'Tampilkan Anomali Usaha dan Keluarga' berhasil diaktifkan.")
                return True
        logger.warning(f"{worker_prefix}Switch anomali belum aktif setelah klik (percobaan {attempt}/3). Mengulang...")

    raise Exception("Gagal mengaktifkan toggle Tampilkan Anomali Usaha dan Keluarga.")

def to_int_or_zero(val):
    """Konversi nilai sel Excel ke integer; kosong/None/non-numerik dianggap 0."""
    try:
        if val is None or str(val).strip() == "":
            return 0
        return int(round(float(val)))
    except (ValueError, TypeError):
        return 0

def parse_numeric_cell(val):
    """Mengembalikan float jika numerik valid. Mengembalikan None jika kosong, None, atau non-numerik."""
    if val is None:
        return None
    s = str(val).strip()
    if s == "":
        return None
    try:
        return float(s)
    except (ValueError, TypeError):
        return None

def fill_components_from_excel(page, comp_list, sheet, row_idx, col_map, event_bus=None, worker_prefix="", excel_lock=None):
    """
    Dipakai saat total BERBEDA (remarks #4):
    - Jika kolom komponen tidak ada di Excel atau bernilai kosong/null: biarkan saja (skip).
    - Jika bernilai 0:
        * Cek apakah elemen komponen ada dan visible di eform Fasih.
        * Jika ada: timpa dengan 0 (agar kalkulasi total di Fasih sinkron dengan Excel).
        * Jika tidak ada: biarkan / lewati (tidak melempar exception).
    - Jika bernilai > 0:
        * Cek apakah elemen komponen ada dan visible di eform Fasih.
        * Jika ada: isi dengan nilai target.
        * Jika tidak ada: lempar exception (karena nilai > 0 wajib terisi).
    """
    for comp in comp_list:
        c_idx = col_map.get(comp["excel_key"])
        if not c_idx:
            continue

        if excel_lock:
            with excel_lock:
                raw_val = sheet.cell(row=row_idx, column=c_idx).value
        else:
            raw_val = sheet.cell(row=row_idx, column=c_idx).value

        num_val = parse_numeric_cell(raw_val)
        if num_val is None:
            # Kosong / null -> biarkan saja
            continue

        target = int(round(num_val))
        loc = page.locator(comp["selector"]).first
        if loc.count() == 0 and "fallback" in comp:
            loc = page.locator(comp["fallback"]).first

        is_present = False
        try:
            is_present = loc.count() > 0 and loc.is_visible()
        except Exception:
            is_present = False

        if target == 0:
            if is_present:
                logger.info(f"{worker_prefix}Mengisi {comp['name']}: 0 (menimpa dengan 0)")
                safe_fill_currency(page, loc, 0, event_bus=event_bus)
                random_sleep(0.3, 0.6, event_bus=event_bus, allow_hesitation=False)
            else:
                logger.debug(f"{worker_prefix}Komponen '{comp['name']}' bernilai 0 di Excel dan elemen tidak ditemukan di form Fasih -> Dilewati.")
        elif target > 0:
            if not is_present:
                raise Exception(f"Field '{comp['name']}' tidak ditemukan di form Fasih, padahal nilai Excel = {target:,}.")

            logger.info(f"{worker_prefix}Mengisi {comp['name']}: {target:,}")
            safe_fill_currency(page, loc, target, event_bus=event_bus)
            random_sleep(0.3, 0.6, event_bus=event_bus, allow_hesitation=False)

def is_values_different(val_clean, val_raw):
    clean_empty = val_clean is None or str(val_clean).strip() == ""
    raw_empty = val_raw is None or str(val_raw).strip() == ""
    if clean_empty and raw_empty:
        return False
    if clean_empty != raw_empty:
        return True
    try:
        clean_num = int(round(float(val_clean)))
        raw_num = int(round(float(val_raw)))
        return clean_num != raw_num
    except (ValueError, TypeError):
        return str(val_clean).strip() != str(val_raw).strip()

def parse_fasih_currency(val_str):
    if not val_str:
        return 0
    cleaned = re.sub(r"[^\d]", "", str(val_str))
    return int(cleaned) if cleaned else 0

def read_container_currency_value(page, selector_or_locator):
    """Membaca nilai nominal mata uang dari elemen container Fasih (melalui input value atau inner_text)."""
    try:
        if isinstance(selector_or_locator, str):
            loc = page.locator(selector_or_locator).first
        else:
            loc = selector_or_locator
        if loc.count() == 0:
            return None
        inp = loc.locator("input[type='text'], input").first
        if inp.count() > 0:
            val = inp.input_value()
            if val is not None and str(val).strip() != "":
                return parse_fasih_currency(val)
        txt = loc.inner_text()
        return parse_fasih_currency(txt)
    except Exception:
        return None

def verify_eform_total(page, selector, target_clean, field_name, worker_prefix="", max_wait_sec=3.5):
    """
    Menunggu dan memverifikasi apakah total di eform Fasih sesuai dengan target_clean dari Excel.
    Jika tidak cocok setelah timeout, melempar Exception sehingga submit dibatalkan.
    """
    target_num = int(round(float(target_clean))) if target_clean is not None and str(target_clean).strip() != "" else 0
    start = time.time()
    last_val = None
    while time.time() - start < max_wait_sec:
        last_val = read_container_currency_value(page, selector)
        if last_val is not None and last_val == target_num:
            logger.bot(f"{worker_prefix}Verifikasi {field_name} COCOK: eform Rp {last_val:,} == Excel Rp {target_num:,}")
            return True
        time.sleep(0.3)

    if last_val is not None and last_val != target_num:
        diff = abs(last_val - target_num)
        err_msg = (
            f"Validasi {field_name} TIDAK COCOK setelah pengisian komponen! "
            f"Total di eform Fasih = Rp {last_val:,}, sedangkan target di Excel (Clean) = Rp {target_num:,} "
            f"(Selisih: Rp {diff:,}). Submit dibatalkan agar data tidak inkonsisten."
        )
        logger.error(f"{worker_prefix}{err_msg}")
        raise Exception(err_msg)
    return True

# ==============================================================================
# ENGINE OTOMASI DATA CLEANING (WORKER RUNNER)
# ==============================================================================

class SessionExpiredException(Exception):
    """Exception khusus saat sesi Fasih SSO kedaluwarsa atau belum terotentikasi di tengah jalan."""
    pass

class DataCleaningEngine:
    def __init__(self, excel_path, event_bus, config_options=None):
        self.excel_path = excel_path
        self.event_bus = event_bus
        self.config_options = config_options or {}
        self.auto_minimize = self.config_options.get("auto_minimize", True)
        self.debug_review_submit = self.config_options.get("debug_review_submit", False)
        self.row_start = self.config_options.get("row_start", 2)
        self.row_end = self.config_options.get("row_end", None)

        # Konfigurasi Mode Login SSO
        self.auth_mode = self.config_options.get("auth_mode", "manual")  # "manual" | "auto"
        self.sso_username = self.config_options.get("sso_username", "").strip()
        self.sso_password = self.config_options.get("sso_password", "").strip()

        # Multi-worker configuration
        raw_workers = self.config_options.get("num_workers", 2)
        try:
            self.num_workers = max(1, min(5, int(raw_workers)))
        except (ValueError, TypeError):
            self.num_workers = 2

        # Jika mode debug review aktif, batasi worker ke 1 agar interaksi review pengguna tidak bentrok
        if self.debug_review_submit and self.num_workers > 1:
            logger.warning("⚠️ Mode Review Debugging aktif: Jumlah worker disetel ke 1 agar review dapat dilakukan teratur.")
            self.num_workers = 1

        # Thread synchronization locks
        self.excel_lock = threading.Lock()
        self.stats_lock = threading.Lock()
        self.reauth_lock = threading.Lock()
        self.task_queue = queue.Queue()

        # Shared session statistics
        self.done_session = 0
        self.error_count = 0
        self.skipped_count = 0
        self.processed_count = 0

    def check_pause_and_stop(self, row_idx=None, audit_data=None, worker_id=None):
        """Memeriksa sinyal jeda atau henti secara responsif di setiap tahapan."""
        if self.event_bus.is_stop_requested():
            if row_idx and audit_data:
                audit_data["status"] = "DIHENTIKAN"
                audit_data["keterangan"] = "Dihentikan oleh pengguna"
                self.event_bus.emit("row_completed", {
                    "row_idx": row_idx,
                    "status": "⏹️ DIHENTIKAN",
                    "keterangan": "Dihentikan oleh pengguna",
                    "worker_id": worker_id
                })
            return False
        self.event_bus.wait_if_paused()
        if self.event_bus.is_stop_requested():
            if row_idx and audit_data:
                audit_data["status"] = "DIHENTIKAN"
                audit_data["keterangan"] = "Dihentikan oleh pengguna"
                self.event_bus.emit("row_completed", {
                    "row_idx": row_idx,
                    "status": "⏹️ DIHENTIKAN",
                    "keterangan": "Dihentikan oleh pengguna",
                    "worker_id": worker_id
                })
            return False
        return True

    def _handle_relogin_needed(self, page, context, row_idx=None, nama_usaha=None, is_midrun=False, worker_id=1):
        """
        Menangani kondisi sesi Fasih SSO kedaluwarsa atau belum login.
        Dikoordinasikan secara thread-safe menggunakan reauth_lock.
        """
        worker_tag = f"[W-{worker_id}] " if worker_id else ""
        with self.reauth_lock:
            # Pengecekan cepat: apakah worker lain baru saja menyelesaikan re-login?
            if is_midrun:
                try:
                    if page and not page.is_closed() and BrowserManager.is_page_logged_in(page) and not BrowserManager.is_session_expired(page):
                        logger.info(f"{worker_tag}Sesi Fasih SSO telah dipulihkan oleh worker lain. Melanjutkan...")
                        return page
                except Exception:
                    pass

            # Opsi 1: Coba Auto-Login SSO BPS jika mode 'auto' dipilih pengguna
            if self.auth_mode == "auto" and self.sso_username and self.sso_password:
                logger.bot(f"{worker_tag}Sesi belum aktif / kedaluwarsa. Menjalankan Auto-Login SSO BPS (User: {self.sso_username})...")
                login_page = page
                if not login_page or login_page.is_closed():
                    if context.pages:
                        login_page = context.pages[0]
                    else:
                        login_page = context.new_page()

                auto_ok, auto_err = BrowserManager.perform_auto_login(
                    login_page, self.sso_username, self.sso_password
                )
                if auto_ok:
                    logger.success(f"🟢 {worker_tag}Auto-Login SSO BPS berhasil! Sesi telah aktif.")
                    if self.auto_minimize:
                        WindowController.minimize_browser()
                    time.sleep(1.5)
                    return login_page
                else:
                    logger.warning(f"⚠️ {worker_tag}Auto-Login gagal ({auto_err}). Beralih ke mode Login Manual...")

            # Opsi 2 / Fallback: Mode Login Manual (meminta pengguna login di Google Chrome)
            if is_midrun and row_idx:
                prefix = f"Baris {row_idx} [{nama_usaha}]"
                msg = f"{worker_tag}Sesi Fasih kedaluwarsa di tengah jalan saat memproses {prefix}. Silakan login kembali di Google Chrome, lalu klik '🟢 Saya Sudah Login'."
                logger.warning(f"⚠️ {worker_tag}{prefix}: Sesi Fasih SSO kedaluwarsa di tengah jalan! Menjeda seluruh worker & menampilkan Google Chrome...")
            else:
                prefix = "Pemeriksaan Sesi Awal"
                msg = "Sesi Fasih SSO belum terdeteksi. Silakan login SSO BPS di Google Chrome, lalu klik '🟢 Saya Sudah Login'."
                logger.warning(f"⚠️ {prefix}: Sesi Fasih SSO belum aktif. Menampilkan jendela Google Chrome...")

            WindowController.restore_and_focus_browser()

            # Reset flag konfirmasi manual sebelumnya
            self.event_bus.reset_auth_confirmation()

            self.event_bus.emit("auth_waiting", {"message": msg, "is_midrun": is_midrun, "row_idx": row_idx, "worker_id": worker_id})
            logger.info(f"{worker_tag}Menunggu Anda login SSO BPS di Google Chrome ({'Mid-Run Baris ' + str(row_idx) if is_midrun else 'Awal'})...")

            active_page = page
            while True:
                # 1. Periksa jika ada permintaan stop dari user
                if self.event_bus.is_stop_requested():
                    logger.warning(f"{worker_tag}Otomasi dihentikan oleh pengguna saat menunggu login SSO.")
                    return None

                # 2. Periksa konfirmasi manual dari tombol UI '🟢 Saya Sudah Login'
                if self.event_bus.is_auth_confirmed():
                    logger.success(f"🟢 {worker_tag}Konfirmasi login diterima dari pengguna! Memeriksa tab browser...")
                    detected = BrowserManager.find_authenticated_page(context)
                    if detected:
                        active_page = detected
                    elif not active_page or active_page.is_closed():
                        active_page = context.pages[0] if context.pages else None
                    break

                # 3. Pindai secara dinamis tab aktif atau tab lain yang terbuka di context (Auto-detect)
                try:
                    if active_page and not active_page.is_closed():
                        if BrowserManager.is_page_logged_in(active_page) and not BrowserManager.is_session_expired(active_page):
                            logger.success(f"✅ {worker_tag}Login SSO Fasih berhasil dideteksi otomatis pada tab aktif!")
                            break
                except Exception:
                    pass

                detected = BrowserManager.find_authenticated_page(context)
                if detected:
                    active_page = detected
                    logger.success(f"✅ {worker_tag}Login SSO Fasih berhasil dideteksi otomatis pada tab browser!")
                    break

                time.sleep(0.5)

            if self.event_bus.is_stop_requested():
                return None

            # Reset flag konfirmasi
            self.event_bus.reset_auth_confirmation()

            # Pastikan active_page masih valid
            if not active_page or active_page.is_closed():
                active_page = BrowserManager.find_authenticated_page(context)
                if not active_page and context.pages:
                    active_page = context.pages[0]

            # KRUSIAL: Beri jeda dan tunggu sampai pengalihan pasca-SSO (SSO Keycloak/OAuth -> Fasih /app) selesai dimuat sepenuhnya
            try:
                active_page.wait_for_load_state("domcontentloaded", timeout=6000)
            except Exception:
                pass
            time.sleep(2.0)

            logger.success(f"🟢 {worker_tag}Sesi Fasih aktif! Melanjutkan proses otomasi...")
            self.event_bus.emit("auth_success", {"is_midrun": is_midrun, "row_idx": row_idx, "worker_id": worker_id})

            if self.auto_minimize:
                WindowController.minimize_browser()

            return active_page

    def _navigate_assignment_page(self, page, raw_url, row_idx=None, nama_usaha="", worker_prefix=""):
        """
        Membuka URL penugasan dengan penanganan tangguh terhadap interupsi navigasi,
        terutama setelah pengalihan SSO Fasih / redirect ke /app.
        """
        max_attempts = 4
        last_err = None
        for attempt in range(max_attempts):
            try:
                # Pastikan tab masih terbuka
                if page.is_closed():
                    raise Exception("Tab browser tertutup sebelum navigasi URL.")

                # Tunggu load state jika ada pengalihan yang sedang berjalan
                try:
                    page.wait_for_load_state("domcontentloaded", timeout=4000)
                except Exception:
                    pass

                page.goto(raw_url, timeout=45000, wait_until="domcontentloaded")
                random_sleep(1.2, 2.0, event_bus=self.event_bus)
                return True
            except Exception as nav_e:
                last_err = nav_e
                nav_err_msg = str(nav_e).lower()
                if any(k in nav_err_msg for k in ["interrupted", "navigation", "abort", "failed to navigate"]):
                    if attempt < max_attempts - 1:
                        logger.warning(
                            f"{worker_prefix}Navigasi baris {row_idx} terinterupsi pengalihan browser ({nav_e}). "
                            f"Menunggu 2.5 detik untuk mencoba ulang ({attempt + 1}/{max_attempts})..."
                        )
                        time.sleep(2.5)
                        continue
                if attempt < max_attempts - 1:
                    time.sleep(1.5)
                else:
                    raise last_err
        return False

    def _process_single_assignment(self, worker_id, page, context, row_idx, wb, sheet, col_map, end_r, target_total):
        """Memproses satu baris penugasan data cleaning dari alur 1 s.d. 8 secara mandiri per worker."""
        w_tag = f"[W-{worker_id}] "

        url_col = col_map.get("url")
        id_col = col_map.get("id")
        name_col = col_map.get("nama_usaha")
        tp_col = col_map.get("total_pengeluaran")
        tp_raw_col = col_map.get("total_pengeluaran_raw")
        tpend_col = col_map.get("total_pendapatan")
        tpend_raw_col = col_map.get("total_pendapatan_raw")
        ta_col = col_map.get("total_aset")
        ta_raw_col = col_map.get("total_aset_raw")
        status_col = col_map.get("status")
        ket_col = col_map.get("keterangan")

        # Baca nilai baris secara aman dari Excel
        with self.excel_lock:
            url_val = sheet.cell(row=row_idx, column=url_col).value if url_col else None
            assignment_id = str(sheet.cell(row=row_idx, column=id_col).value or "").strip() if id_col else ""
            nama_usaha = str(sheet.cell(row=row_idx, column=name_col).value or "").strip() if name_col else ""
            status_val = str(sheet.cell(row=row_idx, column=status_col).value or "").strip() if status_col else ""
            tp_val = sheet.cell(row=row_idx, column=tp_col).value if tp_col else None
            tp_raw_val = sheet.cell(row=row_idx, column=tp_raw_col).value if tp_raw_col else None
            tpend_val = sheet.cell(row=row_idx, column=tpend_col).value if tpend_col else None
            tpend_raw_val = sheet.cell(row=row_idx, column=tpend_raw_col).value if tpend_raw_col else None
            ta_val = sheet.cell(row=row_idx, column=ta_col).value if ta_col else None
            ta_raw_val = sheet.cell(row=row_idx, column=ta_raw_col).value if ta_raw_col else None

        if not url_val or str(url_val).strip() == "":
            return

        with self.stats_lock:
            self.processed_count += 1
            curr_proc = self.processed_count

        # Emit progress update
        self.event_bus.emit("progress_updated", {
            "current": curr_proc,
            "total": target_total,
            "row_idx": row_idx,
            "nama_usaha": nama_usaha,
            "assignment_id": assignment_id,
            "worker_id": worker_id
        })

        if status_val.upper() in ["SUKSES", "DONE"]:
            with self.excel_lock:
                existing_ket = str(sheet.cell(row=row_idx, column=ket_col).value or "").strip() if ket_col else ""
            logger.info(f"{w_tag}Baris {row_idx} [{nama_usaha}] berstatus {status_val} -> Dilewati (keterangan dipertahankan).")
            with self.stats_lock:
                self.skipped_count += 1
            self.event_bus.emit("row_completed", {
                "row_idx": row_idx,
                "status": f"{status_val} (LEWAT)",
                "keterangan": existing_ket,
                "worker_id": worker_id
            })
            return

        # Update Chrome tab visual title dengan identifier nama file & worker
        try:
            file_label = Path(self.excel_path).stem[:10] if self.excel_path else "Fasih"
            page.evaluate(f"() => {{ document.title = '[{file_label}|W-{worker_id}] {nama_usaha[:16]}'; }}")
        except Exception:
            pass

        self.event_bus.emit("worker_status_updated", {
            "worker_id": worker_id,
            "status": "BUSY",
            "row_idx": row_idx,
            "nama_usaha": nama_usaha
        })

        logger.bot(f"{w_tag}--- [Baris {row_idx}/{end_r}] Memproses: {nama_usaha} ---")
        self.event_bus.emit("row_started", {
            "row_idx": row_idx,
            "nama_usaha": nama_usaha,
            "id": assignment_id,
            "worker_id": worker_id
        })

        audit_data = {
            "row_index": row_idx,
            "assignment_id": assignment_id,
            "nama_usaha": nama_usaha,
            "link_fasih": str(url_val).strip(),
            "status": "GAGAL"
        }

        try:
            raw_url = str(url_val).strip()
            if not raw_url.startswith(("http://", "https://")):
                raw_url = "https://" + raw_url

            max_reauth_attempts = 2
            reauth_count = 0

            while reauth_count <= max_reauth_attempts:
                try:
                    # STEP 1: Buka URL Penugasan
                    self._navigate_assignment_page(page, raw_url, row_idx=row_idx, nama_usaha=nama_usaha, worker_prefix=w_tag)

                    # Pengecekan sesi setelah halaman penugasan dimuat
                    if BrowserManager.is_session_expired(page) or not BrowserManager.is_page_logged_in(page):
                        raise SessionExpiredException("Sesi login Fasih tidak aktif atau kedaluwarsa saat membuka formulir.")

                    # STEP 2: Buka Sub Sidebar Blok II / Nama Usaha
                    if not self.check_pause_and_stop(row_idx, audit_data, worker_id):
                        return

                    try:
                        try:
                            page.locator(".fasih-form-sidebar, div[class*='sidebar']").first.wait_for(state="visible", timeout=35000)
                        except Exception:
                            pass

                        try:
                            page.locator(BLOK_II_XPATH).first.wait_for(state="visible", timeout=35000)
                            safe_click(page, BLOK_II_XPATH, timeout=25000, event_bus=self.event_bus)
                        except Exception:
                            safe_click(page, "div[title*='BLOK II'], *[title*='BLOK II']", timeout=20000, event_bus=self.event_bus)
                        random_sleep(0.8, 1.4, event_bus=self.event_bus)
                    except Exception as blok2_err:
                        if BrowserManager.is_session_expired(page) or not BrowserManager.is_page_logged_in(page):
                            raise SessionExpiredException(f"Sesi login Fasih kedaluwarsa saat membuka Blok II: {blok2_err}")
                        raise blok2_err

                    # Cari item nama usaha di sub-sidebar
                    usaha_item = None
                    if nama_usaha:
                        usaha_xpath = f"xpath=//div[contains(@class, 'fasih-form-sidebar')]//div[contains(@class, 'cursor-pointer') and contains(., '{nama_usaha}')] | //div[contains(@class, 'cursor-pointer') and contains(., '{nama_usaha}')]"
                        if page.locator(usaha_xpath).first.count() > 0:
                            usaha_item = page.locator(usaha_xpath).first

                    if not usaha_item:
                        first_sub = page.locator("xpath=//div[contains(@class, 'fasih-form-sidebar')]//div[contains(@class, 'cursor-pointer') and not(contains(., 'BLOK II')) and not(contains(., 'CATATAN'))] | //div[contains(@class, 'cursor-pointer') and not(contains(., 'BLOK II')) and not(contains(., 'CATATAN'))]").first
                        if first_sub.count() > 0:
                            usaha_item = first_sub

                    if usaha_item:
                        safe_click(page, usaha_item, timeout=25000, event_bus=self.event_bus)
                        random_sleep(1.5, 2.5, event_bus=self.event_bus)

                    # Tunggu form utama tampil
                    try:
                        page.locator(TOTAL_PENGELUARAN_SELECTOR).first.wait_for(state="attached", timeout=40000)
                    except Exception as wait_form_err:
                        if BrowserManager.is_session_expired(page) or not BrowserManager.is_page_logged_in(page):
                            raise SessionExpiredException("Sesi Fasih kedaluwarsa saat memuat rincian usaha.")
                        raise Exception("Elemen rincian usaha tidak muncul setelah navigasi.")

                    # STEP 3: Pengeluaran
                    if not self.check_pause_and_stop(row_idx, audit_data, worker_id):
                        return

                    diff_pengeluaran = is_values_different(tp_val, tp_raw_val)
                    audit_data["pengeluaran_clean"] = tp_val
                    audit_data["pengeluaran_raw"] = tp_raw_val

                    if diff_pengeluaran:
                        logger.bot(f"{w_tag}Total Pengeluaran BERBEDA (Clean: {tp_val} vs RAW: {tp_raw_val}) -> Mengisi komponen...")
                        fill_components_from_excel(page, PENGELUARAN_MAP, sheet, row_idx, col_map, event_bus=self.event_bus, worker_prefix=w_tag, excel_lock=self.excel_lock)
                        random_sleep(0.5, 0.8, event_bus=self.event_bus, allow_hesitation=False)
                        verify_eform_total(page, TOTAL_PENGELUARAN_SELECTOR, tp_val, "Total Pengeluaran", worker_prefix=w_tag)
                        remarks_pengeluaran = "#4"
                    else:
                        logger.bot(f"{w_tag}Total Pengeluaran SAMA (Clean: {tp_val} vs RAW: {tp_raw_val}) -> Lewati komponen.")
                        remarks_pengeluaran = "#0"

                    logger.remarks(f"{w_tag}Memberikan remarks '{remarks_pengeluaran}' pada Total Pengeluaran...")
                    add_remarks_to_element(page, TOTAL_PENGELUARAN_SELECTOR, remarks_pengeluaran, event_bus=self.event_bus, worker_prefix=w_tag)
                    audit_data["remarks_pengeluaran"] = remarks_pengeluaran

                    # STEP 4: Pendapatan
                    if not self.check_pause_and_stop(row_idx, audit_data, worker_id):
                        return

                    diff_pendapatan = is_values_different(tpend_val, tpend_raw_val)
                    audit_data["pendapatan_clean"] = tpend_val
                    audit_data["pendapatan_raw"] = tpend_raw_val

                    if diff_pendapatan:
                        logger.bot(f"{w_tag}Total Pendapatan BERBEDA (Clean: {tpend_val} vs RAW: {tpend_raw_val}) -> Mengisi komponen...")
                        fill_components_from_excel(page, PENDAPATAN_MAP, sheet, row_idx, col_map, event_bus=self.event_bus, worker_prefix=w_tag, excel_lock=self.excel_lock)
                        random_sleep(0.5, 0.8, event_bus=self.event_bus, allow_hesitation=False)
                        verify_eform_total(page, TOTAL_PENDAPATAN_SELECTOR, tpend_val, "Total Pendapatan", worker_prefix=w_tag)
                        remarks_pendapatan = "#4"
                    else:
                        logger.bot(f"{w_tag}Total Pendapatan SAMA (Clean: {tpend_val} vs RAW: {tpend_raw_val}) -> Lewati komponen.")
                        remarks_pendapatan = "#0"

                    logger.remarks(f"{w_tag}Memberikan remarks '{remarks_pendapatan}' pada Total Pendapatan...")
                    add_remarks_to_element(page, TOTAL_PENDAPATAN_SELECTOR, remarks_pendapatan, event_bus=self.event_bus, worker_prefix=w_tag)
                    audit_data["remarks_pendapatan"] = remarks_pendapatan

                    # STEP 5: Total Aset
                    if not self.check_pause_and_stop(row_idx, audit_data, worker_id):
                        return

                    diff_aset = is_values_different(ta_val, ta_raw_val)
                    audit_data["aset_clean"] = ta_val
                    audit_data["aset_raw"] = ta_raw_val

                    if diff_aset:
                        tot_aset_num = to_int_or_zero(ta_val)
                        logger.bot(f"{w_tag}Total Aset BERBEDA (Clean: {ta_val} vs RAW: {ta_raw_val}) -> Menghitung target Aset ({tot_aset_num:,})...")

                        elem_tanah = page.locator(ASET_MAP["tanah_bangunan"]["selector"]).first
                        elem_non_tanah = page.locator(ASET_MAP["selain_tanah"]["selector"]).first

                        val_exist_tanah = 0
                        val_exist_lain = 0
                        try:
                            if elem_tanah.count() > 0:
                                inp = elem_tanah.locator("input").first
                                if inp.count() > 0:
                                    val_exist_tanah = parse_fasih_currency(inp.input_value())
                        except Exception:
                            pass

                        try:
                            if elem_non_tanah.count() > 0:
                                inp = elem_non_tanah.locator("input").first
                                if inp.count() > 0:
                                    val_exist_lain = parse_fasih_currency(inp.input_value())
                        except Exception:
                            pass

                        total_exist = val_exist_tanah + val_exist_lain
                        if total_exist > 0:
                            prop_tanah = val_exist_tanah / float(total_exist)
                            aset_tanah_baru = int(round(tot_aset_num * prop_tanah))
                            aset_non_tanah_baru = tot_aset_num - aset_tanah_baru
                        else:
                            pct_tanah = random.randint(60, 70)
                            aset_tanah_baru = int(round(tot_aset_num * (pct_tanah / 100.0)))
                            aset_non_tanah_baru = tot_aset_num - aset_tanah_baru

                        audit_data["aset_tanah_baru"] = aset_tanah_baru
                        audit_data["aset_non_tanah_baru"] = aset_non_tanah_baru

                        if elem_tanah.count() > 0 and elem_tanah.is_visible():
                            safe_fill_currency(page, elem_tanah, aset_tanah_baru, event_bus=self.event_bus)
                            random_sleep(0.3, 0.6, event_bus=self.event_bus, allow_hesitation=False)

                        if elem_non_tanah.count() > 0 and elem_non_tanah.is_visible():
                            safe_fill_currency(page, elem_non_tanah, aset_non_tanah_baru, event_bus=self.event_bus)
                            random_sleep(0.3, 0.6, event_bus=self.event_bus, allow_hesitation=False)

                        random_sleep(0.5, 0.8, event_bus=self.event_bus, allow_hesitation=False)
                        verify_eform_total(page, ASET_MAP["total_aset"]["selector"], ta_val, "Total Aset", worker_prefix=w_tag)
                        remarks_aset = "#4"
                    else:
                        logger.bot(f"{w_tag}Total Aset SAMA (Clean: {ta_val} vs RAW: {ta_raw_val}) -> Lewati pengisian aset.")
                        remarks_aset = "#0"

                    logger.remarks(f"{w_tag}Memberikan remarks '{remarks_aset}' pada Total Aset...")
                    try:
                        add_remarks_to_element(page, ASET_MAP["total_aset"]["selector"], remarks_aset, event_bus=self.event_bus, worker_prefix=w_tag)
                    except Exception:
                        add_remarks_to_element(page, ASET_MAP["tanah_bangunan"]["selector"], remarks_aset, event_bus=self.event_bus, worker_prefix=w_tag)
                    audit_data["remarks_aset"] = remarks_aset

                    # STEP 6: Menuju Blok Catatan di Sidebar
                    if not self.check_pause_and_stop(row_idx, audit_data, worker_id):
                        return

                    try:
                        page.locator(BLOK_CATATAN_XPATH).first.wait_for(state="visible", timeout=35000)
                        safe_click(page, BLOK_CATATAN_XPATH, timeout=25000, event_bus=self.event_bus)
                    except Exception:
                        safe_click(page, "div[title='CATATAN'], *[title='CATATAN']", timeout=25000, event_bus=self.event_bus)
                    random_sleep(1.2, 2.0, event_bus=self.event_bus)

                    # STEP 7: Switch Tampilkan Anomali Usaha dan Keluarga
                    if not self.check_pause_and_stop(row_idx, audit_data, worker_id):
                        return

                    ensure_anomali_switch_active(page, event_bus=self.event_bus, wait_timeout=35000, worker_prefix=w_tag)
                    random_sleep(1.8, 2.5, event_bus=self.event_bus, allow_hesitation=False)

                    # STEP 8: Submit Flow Berantai (4 Langkah)
                    if not self.check_pause_and_stop(row_idx, audit_data, worker_id):
                        return

                    # LANGKAH 1: Klik Kirim pada formulir Fasih
                    logger.bot(f"{w_tag}[1/4] Mengklik tombol Kirim formulir Fasih...")
                    form_kirim_btn = page.locator("button#btn-kirim, button:has-text('Kirim')").first
                    if form_kirim_btn.count() == 0 or not form_kirim_btn.is_visible():
                        form_kirim_btn = page.locator(KIRIM_BUTTON_XPATH).first

                    if form_kirim_btn.count() == 0:
                        raise Exception("Tombol Kirim pada formulir Fasih tidak ditemukan.")

                    safe_click(page, form_kirim_btn, timeout=25000, event_bus=self.event_bus, force=True)
                    random_sleep(1.0, 1.8, event_bus=self.event_bus, allow_hesitation=False)

                    # LANGKAH 2: Dialog Pop-up Validasi Kirim (Pengecekan Jumlah Galat)
                    btn_konfirmasi = page.locator(KONFIRMASI_KIRIM_BUTTON_SELECTOR).last

                    if not btn_konfirmasi.is_visible(timeout=2500):
                        logger.bot(f"{w_tag}[2/4] Menunggu pop-up validasi Kirim...")
                        dialog = page.locator("div[role='dialog'], div[role='alertdialog'], div:has(button:has-text('Batal')):has(button:has-text('Kirim'))").last
                        dialog.wait_for(state="visible", timeout=25000)
                        random_sleep(0.5, 0.8, event_bus=self.event_bus, allow_hesitation=False)

                        # Ambil jumlah Galat dari kartu 'Galat' di dalam pop-up
                        galat_count = 0
                        try:
                            galat_card = dialog.locator("button:has-text('Galat')").first
                            if galat_card.count() > 0 and galat_card.is_visible():
                                card_text = galat_card.inner_text().strip()
                                lines = [l.strip() for l in card_text.split("\n") if l.strip()]
                                for line in lines:
                                    if line.isdigit():
                                        galat_count = int(line)
                                        break
                                else:
                                    nums = re.findall(r"\b\d+\b", card_text)
                                    if nums:
                                        galat_count = int(nums[0])
                            else:
                                span_num = dialog.locator("xpath=.//*[text()='Galat']/following-sibling::*[1]").first
                                if span_num.count() > 0:
                                    t = span_num.inner_text().strip()
                                    if t.isdigit():
                                        galat_count = int(t)
                        except Exception as e_card:
                            logger.bot(f"{w_tag}Catatan: Gagal membaca angka kartu galat: {e_card}")

                        audit_data["galat_count"] = galat_count

                        # HANYA flag ada galat jika jumlah galat > 0
                        if galat_count > 0:
                            err_msg = f"Terdapat {galat_count} Galat yang perlu diperbaiki. Submit dibatalkan."
                            logger.warning(f"⚠️ {w_tag}Baris {row_idx} [{nama_usaha}]: {err_msg}")

                            try:
                                btn_batal = dialog.locator("button:has-text('Batal')").first
                                if btn_batal.is_visible(timeout=3000):
                                    safe_click(page, btn_batal, timeout=5000, event_bus=self.event_bus)
                                else:
                                    page.keyboard.press("Escape")
                            except Exception:
                                page.keyboard.press("Escape")

                            with self.excel_lock:
                                sheet.cell(row=row_idx, column=status_col).value = f"GAGAL ({galat_count} Galat)"
                                sheet.cell(row=row_idx, column=ket_col).value = err_msg
                                try:
                                    wb.save(self.excel_path)
                                except Exception as save_err:
                                    logger.warning(f"{w_tag}Gagal menyimpan Excel: {save_err}")

                            with self.stats_lock:
                                self.error_count += 1

                            audit_data["status"] = "GAGAL"
                            audit_data["keterangan"] = err_msg
                            audit_logger.record_row(audit_data)

                            self.event_bus.emit("row_completed", {
                                "row_idx": row_idx,
                                "status": f"GAGAL ({galat_count} Galat)",
                                "keterangan": err_msg,
                                "worker_id": worker_id
                            })
                            return

                        # Jika Galat == 0, klik tombol Kirim pada dialog
                        logger.bot(f"{w_tag}[2/4] Galat = 0. Mengklik tombol Kirim pada pop-up validasi...")
                        btn_kirim_dialog = dialog.locator("button:has-text('Kirim')").last
                        safe_click(page, btn_kirim_dialog, timeout=20000, event_bus=self.event_bus)
                        random_sleep(1.5, 2.2, event_bus=self.event_bus, allow_hesitation=False)

                    # LANGKAH 3: Dialog Pop-up Konfirmasi Kirim -> Klik Konfirmasi
                    logger.bot(f"{w_tag}[3/4] Menunggu pop-up Konfirmasi Kirim...")
                    btn_konfirmasi = page.locator(KONFIRMASI_KIRIM_BUTTON_SELECTOR).last
                    btn_konfirmasi.wait_for(state="visible", timeout=30000)
                    random_sleep(0.5, 1.0, event_bus=self.event_bus, allow_hesitation=False)

                    if self.debug_review_submit:
                        logger.warning(f"🔍 [MODE REVIEW DEBUGGING] {w_tag}Baris {row_idx} [{nama_usaha}]: Modal Konfirmasi Kirim terbuka.")
                        logger.info("Silakan periksa isian data di browser. Tekan tombol '✅ Konfirmasi Submit & Lanjut' di aplikasi untuk mengirimkan.")
                        WindowController.restore_and_focus_browser()
                        self.event_bus.emit("waiting_confirm_submit", {
                            "row_idx": row_idx,
                            "nama_usaha": nama_usaha,
                            "worker_id": worker_id
                        })
                        self.event_bus.wait_for_submit_confirmation()

                        if self.event_bus.is_stop_requested():
                            logger.warning(f"{w_tag}Proses dihentikan sebelum konfirmasi submit.")
                            audit_data["status"] = "DIHENTIKAN"
                            audit_data["keterangan"] = "Dihentikan oleh pengguna saat review konfirmasi"
                            self.event_bus.emit("row_completed", {
                                "row_idx": row_idx,
                                "status": "⏹️ DIHENTIKAN",
                                "keterangan": "Dihentikan oleh pengguna saat review konfirmasi",
                                "worker_id": worker_id
                            })
                            return

                        logger.bot(f"{w_tag}Konfirmasi diterima dari pengguna. Mengklik tombol Konfirmasi...")
                        safe_click(page, btn_konfirmasi, timeout=20000, event_bus=self.event_bus, force=True)
                        if self.auto_minimize:
                            WindowController.minimize_browser()
                    else:
                        logger.bot(f"{w_tag}[3/4] Dialog Konfirmasi Kirim aktif. Mengklik tombol Konfirmasi...")
                        safe_click(page, btn_konfirmasi, timeout=20000, event_bus=self.event_bus, force=True)

                    # LANGKAH 4: Menunggu Dialog 'Assignment Berhasil Disubmit'
                    logger.bot(f"{w_tag}[4/4] Menunggu server memproses submit (dialog 'Assignment Berhasil Disubmit')...")
                    success_dialog = page.locator(SUCCESS_SUBMIT_DIALOG_SELECTOR).first

                    try:
                        success_dialog.wait_for(state="visible", timeout=50000)
                        logger.success(f"🎉 {w_tag}Baris {row_idx} [{nama_usaha}]: Dialog 'Assignment Berhasil Disubmit' terkonfirmasi!")

                        try:
                            btn_tetap = page.locator("button:has-text('Tetap di Halaman')").first
                            if btn_tetap.is_visible(timeout=2000):
                                safe_click(page, btn_tetap, timeout=5000, event_bus=self.event_bus)
                        except Exception:
                            pass

                        with self.excel_lock:
                            sheet.cell(row=row_idx, column=status_col).value = "SUKSES"
                            sheet.cell(row=row_idx, column=ket_col).value = "Assignment Berhasil Disubmit"
                            try:
                                wb.save(self.excel_path)
                            except Exception as save_err:
                                logger.warning(f"{w_tag}Gagal menyimpan checkpoint ke Excel: {save_err}")

                        with self.stats_lock:
                            self.done_session += 1

                        audit_data["status"] = "SUKSES"
                        audit_data["keterangan"] = "Assignment Berhasil Disubmit"
                        audit_logger.record_row(audit_data)

                        self.event_bus.emit("row_completed", {
                            "row_idx": row_idx,
                            "status": "✅ SUKSES",
                            "keterangan": "Assignment Berhasil Disubmit",
                            "worker_id": worker_id
                        })
                        logger.success(f"{w_tag}✅ Baris {row_idx} [{nama_usaha}] BERHASIL DISUBMIT (SUKSES).")
                        return

                    except Exception as wait_err:
                        err_dialog = page.locator("div[role='alertdialog']:has-text('Gagal'), div:has-text('Gagal mengirimkan'), div:has-text('Terjadi kesalahan')").first
                        if err_dialog.count() > 0 and err_dialog.is_visible():
                            err_text = err_dialog.inner_text().strip().replace("\n", " ")
                            raise Exception(f"Submit ditolak server: {err_text[:120]}")
                        raise Exception(f"Timeout menunggu dialog 'Assignment Berhasil Disubmit' (server belum merespons dalam 50 detik): {wait_err}")

                except SessionExpiredException as sess_err:
                    reauth_count += 1
                    if reauth_count > max_reauth_attempts:
                        logger.error(f"❌ {w_tag}Baris {row_idx} [{nama_usaha}]: Batas maksimum re-login ({max_reauth_attempts}x) tercapai.")
                        raise Exception(f"Sesi Fasih berulang kali kedaluwarsa pada baris ini ({sess_err})")

                    logger.warning(f"⚠️ {w_tag}Baris {row_idx} [{nama_usaha}]: Sesi Fasih SSO kedaluwarsa di tengah jalan! Menjeda & menampilkan Google Chrome...")
                    self.event_bus.request_pause()
                    page = self._handle_relogin_needed(page, context, row_idx=row_idx, nama_usaha=nama_usaha, is_midrun=True, worker_id=worker_id)
                    self.event_bus.request_resume()

                    if not page or self.event_bus.is_stop_requested():
                        audit_data["status"] = "DIHENTIKAN"
                        audit_data["keterangan"] = "Dihentikan oleh pengguna saat sesi login kedaluwarsa"
                        self.event_bus.emit("row_completed", {
                            "row_idx": row_idx,
                            "status": "⏹️ DIHENTIKAN",
                            "keterangan": "Dihentikan oleh pengguna saat sesi login kedaluwarsa",
                            "worker_id": worker_id
                        })
                        return

                    logger.info(f"🔄 {w_tag}Mengulang kembali proses otomasi Baris {row_idx} [{nama_usaha}] dari awal (Percobaan ke-{reauth_count + 1})...")
                    time.sleep(2.0)
                    continue

                except Exception as step_err:
                    if (BrowserManager.is_session_expired(page) or not BrowserManager.is_page_logged_in(page)) and reauth_count < max_reauth_attempts:
                        reauth_count += 1
                        logger.warning(f"⚠️ {w_tag}Baris {row_idx} [{nama_usaha}]: Sesi Fasih terputus di tengah proses ({step_err}). Menjeda & menampilkan Google Chrome...")
                        self.event_bus.request_pause()
                        page = self._handle_relogin_needed(page, context, row_idx=row_idx, nama_usaha=nama_usaha, is_midrun=True, worker_id=worker_id)
                        self.event_bus.request_resume()

                        if not page or self.event_bus.is_stop_requested():
                            audit_data["status"] = "DIHENTIKAN"
                            audit_data["keterangan"] = "Dihentikan oleh pengguna saat sesi login kedaluwarsa"
                            self.event_bus.emit("row_completed", {
                                "row_idx": row_idx,
                                "status": "⏹️ DIHENTIKAN",
                                "keterangan": "Dihentikan oleh pengguna saat sesi login kedaluwarsa",
                                "worker_id": worker_id
                            })
                            return

                        logger.info(f"🔄 {w_tag}Mengulang kembali proses otomasi Baris {row_idx} [{nama_usaha}] dari awal (Percobaan ke-{reauth_count + 1})...")
                        time.sleep(2.0)
                        continue
                    else:
                        raise step_err

        except Exception as row_err:
            error_msg = str(row_err)
            logger.error(f"❌ {w_tag}Baris {row_idx} [{nama_usaha}] GAGAL: {error_msg}")
            with self.excel_lock:
                sheet.cell(row=row_idx, column=status_col).value = "GAGAL"
                sheet.cell(row=row_idx, column=ket_col).value = error_msg[:150]
                try:
                    wb.save(self.excel_path)
                except Exception:
                    pass

            with self.stats_lock:
                self.error_count += 1

            audit_data["status"] = "GAGAL"
            audit_data["keterangan"] = error_msg
            audit_logger.record_row(audit_data)

            self.event_bus.emit("row_completed", {
                "row_idx": row_idx,
                "status": "❌ GAGAL",
                "keterangan": error_msg[:150],
                "worker_id": worker_id
            })

        finally:
            self.event_bus.emit("worker_status_updated", {
                "worker_id": worker_id,
                "status": "IDLE",
                "row_idx": None,
                "nama_usaha": ""
            })
            random_sleep(0.8, 1.5, event_bus=self.event_bus)

    def run(self):
        """Metode utama yang dijalankan di background thread untuk mengoordinasikan Multi-Worker."""
        logger.info(f"Memulai Engine Data Cleaning SE2026 (Konfigurasi Worker: {self.num_workers} Worker)...")
        self.event_bus.set_state(EngineState.RUNNING)

        wb = None

        # 1. Validasi awal template Excel
        val_res = ExcelValidator.validate_excel_file(self.excel_path)
        if not val_res["valid"]:
            logger.error(f"Validasi Excel gagal: {val_res['error']}")
            self.event_bus.set_state(EngineState.IDLE)
            self.event_bus.emit("all_completed", {
                "done": 0, "error": 0, "skipped": 0, "was_stopped": False
            })
            return

        col_map = val_res["col_map"]
        total_data_rows = val_res["total_rows"]

        # Muat workbook untuk di-update per baris
        try:
            wb = openpyxl.load_workbook(self.excel_path)
            sheet = wb.active
        except Exception as e:
            logger.error(f"Gagal memuat workbook Excel: {e}")
            self.event_bus.set_state(EngineState.IDLE)
            self.event_bus.emit("all_completed", {
                "done": 0, "error": 0, "skipped": 0, "was_stopped": False
            })
            return

        # 2. Hubungkan ke Chrome CDP (atau luncurkan otomatis jika belum aktif)
        is_running, b_name = BrowserManager.is_cdp_running()
        if not is_running:
            logger.info("Browser belum aktif di port 9222. Meluncurkan Google Chrome secara otomatis...")
            launched, launch_msg = BrowserManager.launch_browser()
            if not launched:
                logger.error(f"Gagal meluncurkan browser: {launch_msg}")
                self.event_bus.set_state(EngineState.IDLE)
                self.event_bus.emit("all_completed", {
                    "done": 0, "error": 0, "skipped": 0, "was_stopped": False
                })
                wb.close()
                return
            time.sleep(1.5)

        # 3. Pemeriksaan Otentikasi Awal SSO Fasih menggunakan probe ringan
        logger.info("Memeriksa status otentikasi SSO Fasih awal...")
        try:
            with sync_playwright() as p_probe:
                browser_probe = p_probe.chromium.connect_over_cdp(CDP_URL)
                context_probe = browser_probe.contexts[0]
                page_probe = BrowserManager.find_authenticated_page(context_probe)
                if not page_probe:
                    for p_cand in context_probe.pages:
                        if "fasih-sm.bps.go.id" in p_cand.url:
                            page_probe = p_cand
                            break
                    if not page_probe:
                        page_probe = context_probe.pages[0] if context_probe.pages else context_probe.new_page()
                        if "fasih-sm.bps.go.id" not in page_probe.url:
                            try:
                                page_probe.goto("https://fasih-sm.bps.go.id/app", timeout=30000)
                            except Exception:
                                pass

                if not BrowserManager.is_page_logged_in(page_probe) or BrowserManager.is_session_expired(page_probe):
                    page_probe = self._handle_relogin_needed(page_probe, context_probe, is_midrun=False, worker_id=1)
                    if not page_probe or self.event_bus.is_stop_requested():
                        logger.warning("Otomasi dibatalkan sebelum login SSO selesai.")
                        wb.close()
                        self.event_bus.set_state(EngineState.IDLE)
                        self.event_bus.emit("all_completed", {
                            "done": 0, "error": 0, "skipped": 0, "was_stopped": True
                        })
                        return
                elif self.auto_minimize:
                    WindowController.minimize_browser()
        except Exception as e_probe:
            logger.error(f"Gagal memeriksa status SSO Fasih: {e_probe}")
            wb.close()
            self.event_bus.set_state(EngineState.IDLE)
            self.event_bus.emit("all_completed", {
                "done": 0, "error": 0, "skipped": 0, "was_stopped": False
            })
            return

        # 4. Kolom mapping index & persiapan baris
        url_col = col_map.get("url")
        id_col = col_map.get("id")
        name_col = col_map.get("nama_usaha")
        status_col = col_map.get("status")
        ket_col = col_map.get("keterangan")

        with self.excel_lock:
            if not status_col:
                status_col = sheet.max_column + 1
                col_map["status"] = status_col
                sheet.cell(row=1, column=status_col).value = "Status_Eksekusi"
                logger.info(f"Kolom 'Status_Eksekusi' otomatis dibuat di kolom {status_col}.")
            if not ket_col:
                ket_col = sheet.max_column + 1
                col_map["keterangan"] = ket_col
                sheet.cell(row=1, column=ket_col).value = "Keterangan"
                logger.info(f"Kolom 'Keterangan' otomatis dibuat di kolom {ket_col}.")

        max_row = sheet.max_row
        start_r = max(2, self.row_start)
        end_r = min(max_row, self.row_end) if self.row_end else max_row

        # Hitung baris data target yang valid dalam rentang
        target_rows = [
            r for r in range(start_r, end_r + 1)
            if url_col and sheet.cell(row=r, column=url_col).value and str(sheet.cell(row=r, column=url_col).value).strip() != ""
        ]
        target_total = len(target_rows) if target_rows else max(1, end_r - start_r + 1)

        if not target_rows:
            logger.warning("Tidak ditemukan baris penugasan yang valid dalam rentang yang dipilih.")
            wb.close()
            self.event_bus.set_state(EngineState.IDLE)
            self.event_bus.emit("all_completed", {
                "done": 0, "error": 0, "skipped": 0, "was_stopped": False
            })
            return

        actual_workers = min(self.num_workers, len(target_rows))
        logger.info(f"Rentang baris target: {start_r} s.d. {end_r} (Ditemukan {len(target_rows)} baris penugasan). Mengaktifkan {actual_workers} Worker.")

        # Masukkan seluruh target baris ke task_queue
        for r_idx in target_rows:
            self.task_queue.put(r_idx)

        # Worker loop runner function
        def worker_runner(w_id):
            w_tag = f"[W-{w_id}] "
            try:
                with sync_playwright() as p_worker:
                    browser = p_worker.chromium.connect_over_cdp(CDP_URL)
                    context = browser.contexts[0]

                    file_label = Path(self.excel_path).stem[:10] if self.excel_path else "Fasih"

                    # Alokasi halaman tab khusus untuk worker ini
                    # Selalu buat tab baru agar terisolasi sempurna dan tidak saling bajak
                    # antar-worker maupun antar-file/aplikasi yang berjalan paralel
                    page = context.new_page()

                    try:
                        page.evaluate(f"() => {{ document.title = '[{file_label}|W-{w_id}] Fasih 3603 Helper'; }}")
                    except Exception:
                        pass

                    self.event_bus.emit("worker_status_updated", {
                        "worker_id": w_id, "status": "IDLE", "row_idx": None, "nama_usaha": ""
                    })
                    logger.info(f"{w_tag}Tab worker siap beroperasi.")

                    while not self.event_bus.is_stop_requested():
                        self.event_bus.wait_if_paused()
                        if self.event_bus.is_stop_requested():
                            break

                        try:
                            row_idx = self.task_queue.get_nowait()
                        except queue.Empty:
                            break

                        # Self-healing tab check: jika tab tertutup secara tidak sengaja oleh pengguna
                        try:
                            if page.is_closed():
                                logger.warning(f"{w_tag}Tab terdeteksi tertutup! Membuka kembali tab baru secara otomatis...")
                                page = context.new_page()
                                try:
                                    page.evaluate(f"() => {{ document.title = '[{file_label}|W-{w_id}] Fasih 3603 Helper'; }}")
                                except Exception:
                                    pass
                        except Exception as e_reopen:
                            logger.error(f"{w_tag}Gagal memulihkan tab (Google Chrome kemungkinan ditutup): {e_reopen}")
                            self.task_queue.task_done()
                            break

                        try:
                            self._process_single_assignment(
                                worker_id=w_id,
                                page=page,
                                context=context,
                                row_idx=row_idx,
                                wb=wb,
                                sheet=sheet,
                                col_map=col_map,
                                end_r=end_r,
                                target_total=target_total
                            )
                        except Exception as e_row:
                            logger.error(f"{w_tag}Error tak terduga pada baris {row_idx}: {e_row}")
                        finally:
                            self.task_queue.task_done()

                    # Bersihkan tab khusus saat worker selesai agar browser tetap rapi
                    try:
                        page.close()
                    except Exception:
                        pass

                    self.event_bus.emit("worker_status_updated", {
                        "worker_id": w_id, "status": "DONE", "row_idx": None, "nama_usaha": ""
                    })
                    logger.info(f"{w_tag}Worker selesai.")

            except Exception as e_worker:
                logger.error(f"{w_tag}Worker mengalami gangguan koneksi CDP: {e_worker}")
                self.event_bus.emit("worker_status_updated", {
                    "worker_id": w_id, "status": "ERROR", "row_idx": None, "nama_usaha": ""
                })

        try:
            threads = []
            for w_id in range(1, actual_workers + 1):
                t = threading.Thread(target=worker_runner, args=(w_id,), name=f"DataCleaningWorker-{w_id}")
                t.start()
                threads.append(t)
                time.sleep(0.4)

            # Tunggu semua worker selesai
            for t in threads:
                t.join()

        finally:
            if wb:
                try:
                    with self.excel_lock:
                        wb.save(self.excel_path)
                        wb.close()
                except Exception:
                    pass

            if self.auto_minimize:
                try:
                    WindowController.restore_and_focus_browser()
                except Exception:
                    pass

            was_stopped = self.event_bus.is_stop_requested()
            self.event_bus.set_state(EngineState.IDLE)
            if was_stopped:
                logger.warning(f"⏹️ Otomasi Dihentikan oleh Pengguna! (Sukses Sesi Ini: {self.done_session}, Gagal: {self.error_count}, Dilewati: {self.skipped_count})")
            else:
                logger.success(f"🏁 Eksekusi Data Cleaning Selesai! (Sukses Sesi Ini: {self.done_session}, Gagal: {self.error_count}, Dilewati: {self.skipped_count})")

            self.event_bus.emit("all_completed", {
                "done": self.done_session,
                "error": self.error_count,
                "skipped": self.skipped_count,
                "was_stopped": was_stopped
            })

CleaningUsahaEngine = DataCleaningEngine
