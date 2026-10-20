import os
from pathlib import Path
import openpyxl
from modules.data_cleaning.selectors import EXCEL_COLUMNS_CONFIG

class ExcelValidator:
    MANDATORY_KEYS = [
        "url",
        "total_pengeluaran",
        "total_pengeluaran_raw",
        "total_pendapatan",
        "total_pendapatan_raw",
        "total_aset",
        "total_aset_raw"
    ]

    @staticmethod
    def is_file_locked(file_path):
        """Memeriksa apakah file Excel sedang terbuka di aplikasi lain (Microsoft Excel) yang mengunci penulisan."""
        if not os.path.exists(file_path):
            return False
        try:
            with open(file_path, "r+b"):
                return False
        except (IOError, PermissionError):
            return True

    @classmethod
    def validate_excel_file(cls, file_path):
        """
        Melakukan validasi menyeluruh terhadap file Excel template Data Cleaning.
        Mengembalikan dict hasil pemeriksaan.
        """
        p = Path(file_path)
        if not p.exists():
            return {
                "valid": False,
                "error": f"File tidak ditemukan: {file_path}",
                "warning": None,
                "file_locked": False,
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        locked = cls.is_file_locked(file_path)
        warning_msg = None
        if locked:
            warning_msg = "Perhatian: File sedang dibuka di Excel. Anda dapat melihat pratinjau, namun harap tutup file Excel sebelum memulai otomasi agar hasil dapat disimpan."

        try:
            with open(file_path, "rb") as f:
                wb = openpyxl.load_workbook(f, data_only=True)
                sheet = wb.active
        except Exception as e:
            return {
                "valid": False,
                "error": f"Gagal membaca format file Excel: {e}",
                "warning": None,
                "file_locked": locked,
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        # Baca header baris 1
        header_map = {}
        for col_idx in range(1, sheet.max_column + 1):
            val = sheet.cell(row=1, column=col_idx).value
            if val is not None:
                header_map[str(val).strip().lower()] = col_idx

        # Mapping konfigurasi ke index kolom sebenarnya
        col_map = {}
        for key, aliases in EXCEL_COLUMNS_CONFIG.items():
            for alias in aliases:
                clean_alias = str(alias).strip().lower()
                if clean_alias in header_map:
                    col_map[key] = header_map[clean_alias]
                    break

        # Periksa kolom wajib
        missing_mandatory = []
        for mand in cls.MANDATORY_KEYS:
            if mand not in col_map:
                missing_mandatory.append(mand.replace("_", " ").title())

        if missing_mandatory:
            wb.close()
            return {
                "valid": False,
                "error": f"Template tidak sesuai. Kolom wajib berikut tidak ditemukan: {', '.join(missing_mandatory)}",
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        # Parsing baris data untuk preview & counting
        rows_data = []
        total_rows = 0
        done_rows = 0
        pending_rows = 0

        url_col = col_map.get("url")
        name_col = col_map.get("nama_usaha")
        id_col = col_map.get("id")
        tp_col = col_map.get("total_pengeluaran")
        tp_raw_col = col_map.get("total_pengeluaran_raw")
        tpend_col = col_map.get("total_pendapatan")
        tpend_raw_col = col_map.get("total_pendapatan_raw")
        ta_col = col_map.get("total_aset")
        ta_raw_col = col_map.get("total_aset_raw")
        status_col = col_map.get("status")
        ket_col = col_map.get("keterangan")

        for r in range(2, sheet.max_row + 1):
            url_val = sheet.cell(row=r, column=url_col).value if url_col else None
            # Jika kolom URL kosong, abaikan jika seluruh baris kosong
            if not url_val or str(url_val).strip() == "":
                continue

            total_rows += 1
            st_val = str(sheet.cell(row=r, column=status_col).value or "").strip() if status_col else ""
            is_done = st_val.upper() in ["SUKSES", "DONE"]
            if is_done:
                done_rows += 1
            else:
                pending_rows += 1

            rows_data.append({
                "row_idx": r,
                "id": str(sheet.cell(row=r, column=id_col).value or "") if id_col else "",
                "nama_usaha": str(sheet.cell(row=r, column=name_col).value or "") if name_col else "",
                "url": str(url_val).strip(),
                "pengeluaran": sheet.cell(row=r, column=tp_col).value if tp_col else None,
                "pengeluaran_raw": sheet.cell(row=r, column=tp_raw_col).value if tp_raw_col else None,
                "pendapatan": sheet.cell(row=r, column=tpend_col).value if tpend_col else None,
                "pendapatan_raw": sheet.cell(row=r, column=tpend_raw_col).value if tpend_raw_col else None,
                "aset": sheet.cell(row=r, column=ta_col).value if ta_col else None,
                "aset_raw": sheet.cell(row=r, column=ta_raw_col).value if ta_raw_col else None,
                "status": st_val if st_val else "BELUM",
                "keterangan": str(sheet.cell(row=r, column=ket_col).value or "") if ket_col else ""
            })

        wb.close()

        if total_rows == 0:
            return {
                "valid": False,
                "error": "File Excel tidak memiliki baris data penugasan (kolom link_fasih_sm kosong).",
                "total_rows": 0,
                "done_rows": 0,
                "pending_rows": 0,
                "rows_data": []
            }

        return {
            "valid": True,
            "error": None,
            "warning": warning_msg,
            "file_locked": locked,
            "col_map": col_map,
            "total_rows": total_rows,
            "done_rows": done_rows,
            "pending_rows": pending_rows,
            "rows_data": rows_data
        }

    @classmethod
    def generate_template(cls, output_path):
        """
        Membuat file template Excel standar Data Cleaning (template_flag_clean.xlsx)
        lengkap dengan seluruh kolom resmi dan 2 baris contoh data.
        """
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "flag_clean"

        headers = [
            "No", "assignment_id", "Kategori", "KBLI", "KBLI 2 DGT", "Nama Usaha",
            "Jumlah Tenaga Kerja", "Upah dan Gaji", "Biaya Produksi",
            "Biaya Pembelian Barang yang Terjual", "Biaya Operasional", "Biaya Non Operasional",
            "Total Pengeluaran", "Total Pengeluaran [RAW]",
            "Nilai Pendapatan Barang dan Jasa", "Pendapatan Lainnya",
            "Total Pendapatan", "Total Pendapatan [RAW]",
            "Output", "Nilai Tambah",
            "Total Aset", "Total Aset [RAW]",
            "Produktivitas", "Rasio NTB", "Rasio Output terhadap Aset", "Rasio Upah terhadap NTB",
            "Flag 1", "Flag 2", "Flag 3",
            "link_fasih_sm",
            "is_remark_sesuai_lapangan", "is_remark_diperbaiki",
            "Nama Petugas DC", "Penanganan DC\n(Sudah/Belum)",
            "Status_Eksekusi", "Keterangan"
        ]

        ws.append(headers)

        # Style header
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='D1D5DB'),
            right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'),
            bottom=Side(style='thin', color='D1D5DB')
        )

        ws.row_dimensions[1].height = 28

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border

        # 2 Baris Contoh Data
        sample_rows = [
            [
                1, "3603010001001", "C", "10710", "10", "CV Contoh Berkah Mandiri",
                5, 25000000, 15000000, 0, 5000000, 5000000,
                50000000, 40000000,
                75000000, 5000000,
                80000000, 60000000,
                80000000, 30000000,
                120000000, 100000000,
                16000000, 0.375, 0.667, 0.833,
                1, 0, 0,
                "https://fasih-sm.bps.go.id/penugasan/detail/contoh-id-1",
                1, 1, "Admin DC", "Sudah",
                None, None
            ],
            [
                2, "3603010001002", "G", "47111", "47", "Toko Sejahtera Bersama",
                3, 12000000, 0, 30000000, 3000000, 0,
                45000000, 45000000,
                60000000, 0,
                60000000, 60000000,
                60000000, 15000000,
                50000000, 50000000,
                20000000, 0.250, 1.200, 0.800,
                0, 0, 0,
                "https://fasih-sm.bps.go.id/penugasan/detail/contoh-id-2",
                1, 0, "Admin DC", "Sudah",
                None, None
            ]
        ]

        data_font = Font(name="Calibri", size=10)
        data_align_left = Alignment(horizontal="left", vertical="center")
        data_align_right = Alignment(horizontal="right", vertical="center")
        data_align_center = Alignment(horizontal="center", vertical="center")

        for row_data in sample_rows:
            ws.append(row_data)
            current_r = ws.max_row
            ws.row_dimensions[current_r].height = 20
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=current_r, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                val = cell.value
                if isinstance(val, (int, float)):
                    if val > 1000:
                        cell.number_format = "#,##0"
                    cell.alignment = data_align_right
                elif col_idx in [1, 27, 28, 29, 31, 32]:
                    cell.alignment = data_align_center
                else:
                    cell.alignment = data_align_left

        # Atur lebar kolom otomatis
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val_str = str(cell.value or "")
                if "\n" in val_str:
                    val_str = max(val_str.split("\n"), key=len)
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(output_path)
        wb.close()
        return True
