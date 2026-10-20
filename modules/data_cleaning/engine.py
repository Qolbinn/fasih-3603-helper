import os
import re
import time
import random
import openpyxl
from datetime import datetime
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

from config import CDP_URL
from core.logger import logger
from core.event_bus import EngineState
from core.audit_logger import audit_logger
from core.window_controller import WindowController
from modules.data_cleaning.selectors import (
    PENGELUARAN_MAP,
    PENDAPATAN_MAP,
    ASET_MAP,
    TOTAL_PENGELUARAN_SELECTOR,
    TOTAL_PENDAPATAN_SELECTOR,
    BLOK_II_XPATH,
    BLOK_CATATAN_XPATH,
    KIRIM_BUTTON_XPATH,
    CONFIRM_BUTTON_XPATH,
    RINGKASAN_KIRIM_SELECTOR,
    KONFIRMASI_KIRIM_BUTTON_SELECTOR,
    SUCCESS_SUBMIT_DIALOG_SELECTOR
)
from modules.data_cleaning.validator import ExcelValidator

# ==============================================================================
# HELPER BEHAVIOR MANUSIA & PLAYWRIGHT
# ==============================================================================

def responsive_sleep(min_sec=1.0, max_sec=2.0, event_bus=None, allow_hesitation=True):
    duration = random.uniform(min_sec, max_sec)
    if allow_hesitation and random.random() < 0.2:
        duration += random.uniform(1.0, 2.0)
    end_time = time.time() + duration
    while time.time() < end_time:
        if event_bus and event_bus.is_stop_requested():
            return False
        if event_bus and event_bus.is_paused():
            event_bus.wait_if_paused()
            if event_bus.is_stop_requested():
                return False
        time.sleep(min(0.08, max(0.01, end_time - time.time())))
    return True

def human_jitter(min_sec=0.1, max_sec=0.25, event_bus=None):
    responsive_sleep(min_sec, max_sec, event_bus=event_bus, allow_hesitation=False)

def random_sleep(min_sec=1.0, max_sec=2.0, event_bus=None, allow_hesitation=True):
    responsive_sleep(min_sec, max_sec, event_bus=event_bus, allow_hesitation=allow_hesitation)

def safe_click(page, selector_or_locator, timeout=10000, event_bus=None, force=False):
    human_jitter(0.1, 0.2, event_bus=event_bus)
    if isinstance(selector_or_locator, str):
        loc = page.locator(selector_or_locator).first
    else:
        loc = selector_or_locator
    loc.wait_for(state="attached", timeout=timeout)
    try:
        loc.scroll_into_view_if_needed(timeout=3000)
    except Exception:
        pass
    human_jitter(0.1, 0.2, event_bus=event_bus)
    try:
        loc.click(timeout=timeout, force=force)
    except Exception as e:
        err_msg = str(e).lower()
        if "intercepts pointer events" in err_msg or "subtree intercepts" in err_msg or not force:
            logger.bot("Intersepsi klik terdeteksi (animasi/backdrop modal). Mengulang klik dengan force=True...")
            human_jitter(0.1, 0.2, event_bus=event_bus)
            loc.click(timeout=5000, force=True)
        else:
            raise e

def safe_fill_currency(page, container_locator, value, event_bus=None):
    val_int = int(round(float(value)))
    val_str = str(val_int)
    container_locator.scroll_into_view_if_needed()
    input_elem = container_locator.locator("input[type='text'], input").first
    input_elem.wait_for(state="visible", timeout=10000)
    input_elem.click()
    human_jitter(0.1, 0.2, event_bus=event_bus)
    page.keyboard.press("Control+A")
    human_jitter(0.05, 0.1, event_bus=event_bus)
    page.keyboard.press("Backspace")
    human_jitter(0.05, 0.1, event_bus=event_bus)
    input_elem.fill(val_str)
    human_jitter(0.1, 0.2, event_bus=event_bus)
    page.keyboard.press("Tab")

def add_remarks_to_element(page, container_selector, note_text, event_bus=None):
    """Membuka modal Catatan pada elemen container, mengetik catatan, klik simpan, lalu tutup modal."""
    container = page.locator(container_selector).first
    container.wait_for(state="attached", timeout=15000)
    container.scroll_into_view_if_needed()
    random_sleep(0.4, 0.7, event_bus=event_bus)

    # Cari tombol Catatan yang VISIBLE di dalam container
    btn_catatan = container.locator("button[title='Catatan']:visible, button[aria-label='Catatan']:visible").first
    if btn_catatan.count() == 0 or not btn_catatan.is_visible():
        all_btns = container.locator("button[title='Catatan'], button[aria-label='Catatan']")
        btn_catatan = None
        for i in range(all_btns.count()):
            b = all_btns.nth(i)
            if b.is_visible():
                btn_catatan = b
                break

    if not btn_catatan:
        raise Exception(f"Tombol Catatan visible tidak ditemukan pada container {container_selector}")

    safe_click(page, btn_catatan, timeout=10000, event_bus=event_bus)
    random_sleep(0.7, 1.1, event_bus=event_bus)

    dialog = page.locator("div[role='dialog']").first
    try:
        dialog.wait_for(state="visible", timeout=10000)
    except PlaywrightTimeoutError:
        raise Exception("Modal pop-up Catatan tidak muncul setelah tombol Catatan diklik")

    note_str = str(note_text).strip()

    # Cek apakah catatan sudah pernah tercatat sebelumnya di riwayat pesan
    already_exists = False
    try:
        bubbles = dialog.locator("div[class*='whitespace-pre-wrap'], div[class*='wrap-break-word']")
        for i in range(bubbles.count()):
            if bubbles.nth(i).inner_text().strip() == note_str:
                already_exists = True
                break
    except Exception:
        pass

    if already_exists:
        logger.info(f"Catatan '{note_str}' sudah tercatat sebelumnya. Melewati pengiriman ulang.")
    else:
        textarea = dialog.locator("textarea[placeholder*='Tambah catatan'], textarea").first
        try:
            textarea.wait_for(state="visible", timeout=10000)
        except PlaywrightTimeoutError:
            raise Exception("Textarea Catatan tidak ditemukan di dalam dialog")

        textarea.click()
        human_jitter(0.1, 0.2, event_bus=event_bus)
        textarea.fill(note_str)
        human_jitter(0.15, 0.25, event_bus=event_bus)

        btn_simpan = dialog.locator("button[title='Simpan']").first
        btn_simpan.wait_for(state="visible", timeout=10000)
        safe_click(page, btn_simpan, timeout=10000, event_bus=event_bus)
        random_sleep(0.8, 1.2, event_bus=event_bus)

    try:
        btn_close = dialog.locator("button[aria-label='Dismiss']").first
        if btn_close.is_visible(timeout=3000):
            safe_click(page, btn_close, timeout=5000, event_bus=event_bus)
        else:
            page.keyboard.press("Escape")
    except Exception:
        page.keyboard.press("Escape")

    try:
        dialog.wait_for(state="hidden", timeout=5000)
    except Exception:
        page.keyboard.press("Escape")

    random_sleep(0.4, 0.7, event_bus=event_bus)

def ensure_anomali_switch_active(page, event_bus=None, wait_timeout=15000):
    """
    Pastikan switch 'Tampilkan Anomali Usaha dan Keluarga' (div#cek_anomali_button) AKTIF.
    - Sudah aktif (input aria-checked='true')  -> dibiarkan, tidak diklik.
    - Belum aktif (aria-checked='false')       -> klik kontrol switch, lalu verifikasi benar-benar aktif.
    Raise Exception jika elemen tidak muncul atau gagal diaktifkan.
    """
    box = page.locator("div#cek_anomali_button").first
    try:
        # Blok Catatan baru dibuka -> tunggu elemen benar-benar dirender (bukan sekadar count() == 0 lalu dilewati)
        box.wait_for(state="visible", timeout=wait_timeout)
    except PlaywrightTimeoutError:
        raise Exception("Switch 'Tampilkan Anomali Usaha dan Keluarga' tidak muncul di Blok Catatan.")

    box.scroll_into_view_if_needed()
    sw_input = box.locator("input[role='switch']").first
    sw_control = box.locator("[id^='switch-'][id$='-control']").first

    def _is_active():
        return sw_input.get_attribute("aria-checked") == "true"

    if _is_active():
        logger.bot("Switch 'Tampilkan Anomali Usaha dan Keluarga' sudah aktif -> Dibiarkan.")
        return False

    logger.bot("Switch 'Tampilkan Anomali Usaha dan Keluarga' belum aktif -> Mengaktifkan...")
    for attempt in range(1, 4):
        # Input asli disembunyikan (1px, clip) -> yang diklik adalah div kontrol switch-nya
        safe_click(page, sw_control, timeout=5000, event_bus=event_bus)

        # Tunggu state berubah (maks ~3 detik)
        for _ in range(10):
            random_sleep(0.25, 0.35, event_bus=event_bus, allow_hesitation=False)
            if _is_active():
                logger.bot("Switch 'Tampilkan Anomali Usaha dan Keluarga' berhasil diaktifkan.")
                return True
        logger.warning(f"Switch anomali belum aktif setelah klik (percobaan {attempt}/3). Mengulang...")

    raise Exception("Gagal mengaktifkan toggle Tampilkan Anomali Usaha dan Keluarga.")

def to_int_or_zero(val):
    """Konversi nilai sel Excel ke integer; kosong/None/non-numerik dianggap 0."""
    try:
        if val is None or str(val).strip() == "":
            return 0
        return int(round(float(val)))
    except (ValueError, TypeError):
        return 0

def fill_components_from_excel(page, comp_list, sheet, row_idx, col_map, event_bus=None):
    """
    Dipakai saat total BERBEDA (remarks #4): Mengisi komponen penyusun HANYA jika kolomnya ada di Excel
    dan memiliki nilai valid > 0 (sesuai data_cleaning.py & alur.md).
    Jika kolom tidak ada di Excel atau bernilai kosong/0, komponen TIDAK ditimpa dengan 0.
    """
    for comp in comp_list:
        c_idx = col_map.get(comp["excel_key"])
        if not c_idx:
            continue

        raw_val = sheet.cell(row=row_idx, column=c_idx).value
        if not is_valid_positive_number(raw_val):
            continue

        target = int(round(float(raw_val)))
        loc = page.locator(comp["selector"]).first
        if loc.count() == 0 and "fallback" in comp:
            loc = page.locator(comp["fallback"]).first

        if loc.count() == 0 or not loc.is_visible():
            raise Exception(f"Field '{comp['name']}' tidak ditemukan di form Fasih, padahal nilai Excel = {target:,}.")

        logger.info(f"Mengisi {comp['name']}: {target:,}")
        safe_fill_currency(page, loc, target, event_bus=event_bus)
        random_sleep(0.3, 0.6, event_bus=event_bus, allow_hesitation=False)

def is_valid_positive_number(val):
    if val is None:
        return False
    try:
        num = float(val)
        return num > 0
    except (ValueError, TypeError):
        return False

def is_values_different(val_clean, val_raw):
    clean_empty = val_clean is None or str(val_clean).strip() == ""
    raw_empty = val_raw is None or str(val_raw).strip() == ""
    if clean_empty and raw_empty:
        return False
    if clean_empty != raw_empty:
        return True
    try:
        clean_num = int(round(float(val_clean)))
        raw_num = int(round(float(val_raw)))
        return clean_num != raw_num
    except (ValueError, TypeError):
        return str(val_clean).strip() != str(val_raw).strip()

def parse_fasih_currency(val_str):
    if not val_str:
        return 0
    cleaned = re.sub(r"[^\d]", "", str(val_str))
    return int(cleaned) if cleaned else 0

# ==============================================================================
# ENGINE OTOMASI DATA CLEANING (WORKER RUNNER)
# ==============================================================================

class DataCleaningEngine:
    def __init__(self, excel_path, event_bus, config_options=None):
        self.excel_path = excel_path
        self.event_bus = event_bus
        self.config_options = config_options or {}
        self.auto_minimize = self.config_options.get("auto_minimize", True)
        self.debug_review_submit = self.config_options.get("debug_review_submit", False)
        self.row_start = self.config_options.get("row_start", 2)
        self.row_end = self.config_options.get("row_end", None)

    def check_pause_and_stop(self, row_idx=None, audit_data=None):
        """Memeriksa sinyal jeda atau henti secara responsif di setiap tahapan."""
        if self.event_bus.is_stop_requested():
            if row_idx and audit_data:
                audit_data["status"] = "DIHENTIKAN"
                audit_data["keterangan"] = "Dihentikan oleh pengguna"
                self.event_bus.emit("row_completed", {
                    "row_idx": row_idx,
                    "status": "⏹️ DIHENTIKAN",
                    "keterangan": "Dihentikan oleh pengguna"
                })
            return False
        self.event_bus.wait_if_paused()
        if self.event_bus.is_stop_requested():
            if row_idx and audit_data:
                audit_data["status"] = "DIHENTIKAN"
                audit_data["keterangan"] = "Dihentikan oleh pengguna"
                self.event_bus.emit("row_completed", {
                    "row_idx": row_idx,
                    "status": "⏹️ DIHENTIKAN",
                    "keterangan": "Dihentikan oleh pengguna"
                })
            return False
        return True

    def _handle_relogin_needed(self, page, context, row_idx=None, nama_usaha=None):
        """
        Menangani kondisi sesi Fasih SSO kedaluwarsa atau belum login.
        Menampilkan browser ke depan, memancarkan event auth_waiting,
        dan menunggu pengguna login kembali (didukung auto-detect atau klik tombol konfirmasi).
        Mengembalikan tab aktif yang sudah terotentikasi, atau None jika dibatalkan/dihentikan.
        """
        prefix = f"Baris {row_idx} [{nama_usaha}]" if row_idx else "Pemeriksaan Sesi Awal"
        logger.warning(f"⚠️ {prefix}: Sesi Fasih SSO kedaluwarsa / belum aktif. Menampilkan jendela Chrome...")
        WindowController.restore_and_focus_browser()

        # Reset flag konfirmasi manual sebelumnya
        self.event_bus.reset_auth_confirmation()

        msg = f"Sesi Fasih kedaluwarsa saat memproses {prefix}. Silakan login SSO BPS & OTP di Google Chrome (atau klik '🟢 Saya Sudah Login')."
        self.event_bus.emit("auth_waiting", {"message": msg})
        logger.info("Menunggu Anda menyelesaikan login SSO BPS & OTP di browser Chrome...")

        active_page = page
        while True:
            # 1. Periksa jika ada permintaan stop dari user
            if self.event_bus.is_stop_requested():
                logger.warning("Otomasi dihentikan oleh pengguna saat menunggu login SSO.")
                return None

            # 2. Periksa konfirmasi manual dari tombol UI '🟢 Saya Sudah Login'
            if self.event_bus.is_auth_confirmed():
                logger.success("🟢 Konfirmasi login diterima dari tombol UI. Memeriksa tab browser...")
                detected = BrowserManager.find_authenticated_page(context)
                if detected:
                    active_page = detected
                break

            # 3. Pindai secara dinamis tab aktif atau tab lain yang terbuka di context
            try:
                if active_page and not active_page.is_closed():
                    if BrowserManager.is_page_logged_in(active_page) and not BrowserManager.is_session_expired(active_page):
                        logger.success("✅ Login SSO Fasih berhasil dideteksi otomatis pada tab aktif!")
                        break
            except Exception:
                pass

            detected = BrowserManager.find_authenticated_page(context)
            if detected:
                active_page = detected
                logger.success("✅ Login SSO Fasih berhasil dideteksi otomatis pada tab browser!")
                break

            time.sleep(1.0)

        if self.event_bus.is_stop_requested():
            return None

        # Stabilisasi redirect pasca login SSO
        logger.info("Menunggu browser menyelesaikan stabilisasi sesi login SSO...")
        for _ in range(20):
            if self.event_bus.is_stop_requested():
                return None
            try:
                if active_page and not active_page.is_closed():
                    curr_url = active_page.url.lower()
                    if "oauth2" not in curr_url and "sso.bps.go.id" not in curr_url and not BrowserManager.is_session_expired(active_page):
                        break
            except Exception:
                pass
            time.sleep(0.5)

        try:
            if active_page and not active_page.is_closed():
                active_page.wait_for_load_state("domcontentloaded", timeout=5000)
        except Exception:
            pass
        time.sleep(1.0)

        self.event_bus.emit("auth_success", {})

        if self.auto_minimize:
            WindowController.minimize_browser()

        return active_page

    def run(self):
        """Metode utama yang dijalankan di background worker thread."""
        logger.info(f"Memulai Engine Data Cleaning SE2026...")
        self.event_bus.set_state(EngineState.RUNNING)

        wb = None
        p = None
        browser = None
        done_session = 0
        error_count = 0
        skipped_count = 0

        # 1. Validasi awal template Excel
        val_res = ExcelValidator.validate_excel_file(self.excel_path)
        if not val_res["valid"]:
            logger.error(f"Validasi Excel gagal: {val_res['error']}")
            self.event_bus.set_state(EngineState.IDLE)
            self.event_bus.emit("all_completed", {
                "done": 0, "error": 0, "skipped": 0, "was_stopped": False
            })
            return

        col_map = val_res["col_map"]
        total_data_rows = val_res["total_rows"]

        # Muat workbook untuk di-update per baris
        try:
            wb = openpyxl.load_workbook(self.excel_path)
            sheet = wb.active
        except Exception as e:
            logger.error(f"Gagal memuat workbook Excel: {e}")
            self.event_bus.set_state(EngineState.IDLE)
            self.event_bus.emit("all_completed", {
                "done": 0, "error": 0, "skipped": 0, "was_stopped": False
            })
            return

        # 2. Hubungkan ke Chrome CDP (atau luncurkan otomatis jika belum aktif)
        from core.browser_manager import BrowserManager
        is_running, b_name = BrowserManager.is_cdp_running()
        if not is_running:
            logger.info("Browser belum aktif di port 9222. Meluncurkan Google Chrome secara otomatis...")
            launched, launch_msg = BrowserManager.launch_browser()
            if not launched:
                logger.error(f"Gagal meluncurkan browser: {launch_msg}")
                self.event_bus.set_state(EngineState.IDLE)
                self.event_bus.emit("all_completed", {
                    "done": 0, "error": 0, "skipped": 0, "was_stopped": False
                })
                wb.close()
                return

        logger.info(f"Menghubungkan ke browser via CDP ({CDP_URL})...")
        p = sync_playwright().start()
        try:
            browser = p.chromium.connect_over_cdp(CDP_URL)
        except Exception as e:
            logger.error(f"Gagal terhubung ke port CDP Chrome: {e}. Pastikan Chrome sudah dibuka dengan port 9222.")
            p.stop()
            wb.close()
            self.event_bus.set_state(EngineState.IDLE)
            self.event_bus.emit("all_completed", {
                "done": 0, "error": 0, "skipped": 0, "was_stopped": False
            })
            return

        context = browser.contexts[0]

        # 3. Pemeriksaan Otentikasi SSO Fasih (Multi-tab scanning)
        page = BrowserManager.find_authenticated_page(context)

        if not page:
            # Cari tab Fasih yang ada atau buka baru
            for p_cand in context.pages:
                if "fasih-sm.bps.go.id" in p_cand.url:
                    page = p_cand
                    break

            if not page:
                page = context.pages[0] if context.pages else context.new_page()
                if "fasih-sm.bps.go.id" not in page.url:
                    try:
                        page.goto("https://fasih-sm.bps.go.id/app", timeout=30000)
                    except Exception:
                        pass

        # Jika masih belum terotentikasi atau sesi habis, masuki sesi tunggu login
        if not BrowserManager.is_page_logged_in(page) or BrowserManager.is_session_expired(page):
            page = self._handle_relogin_needed(page, context)
            if not page or self.event_bus.is_stop_requested():
                logger.warning("Otomasi dibatalkan sebelum login SSO selesai.")
                browser.close()
                p.stop()
                wb.close()
                self.event_bus.set_state(EngineState.IDLE)
                self.event_bus.emit("all_completed", {
                    "done": 0, "error": 0, "skipped": 0, "was_stopped": True
                })
                return
        elif self.auto_minimize:
            WindowController.minimize_browser()

        # Kolom mapping index
        url_col = col_map.get("url")
        id_col = col_map.get("id")
        name_col = col_map.get("nama_usaha")
        tp_col = col_map.get("total_pengeluaran")
        tp_raw_col = col_map.get("total_pengeluaran_raw")
        tpend_col = col_map.get("total_pendapatan")
        tpend_raw_col = col_map.get("total_pendapatan_raw")
        ta_col = col_map.get("total_aset")
        ta_raw_col = col_map.get("total_aset_raw")
        status_col = col_map.get("status")
        ket_col = col_map.get("keterangan")

        # Auto-create kolom Status_Eksekusi dan Keterangan jika belum ada di template input
        if not status_col:
            status_col = sheet.max_column + 1
            sheet.cell(row=1, column=status_col).value = "Status_Eksekusi"
            logger.info(f"Kolom 'Status_Eksekusi' otomatis dibuat di kolom {status_col}.")
        if not ket_col:
            ket_col = sheet.max_column + 1
            sheet.cell(row=1, column=ket_col).value = "Keterangan"
            logger.info(f"Kolom 'Keterangan' otomatis dibuat di kolom {ket_col}.")

        max_row = sheet.max_row
        start_r = max(2, self.row_start)
        end_r = min(max_row, self.row_end) if self.row_end else max_row

        # Hitung baris data target yang valid dalam rentang
        target_rows = [
            r for r in range(start_r, end_r + 1)
            if url_col and sheet.cell(row=r, column=url_col).value and str(sheet.cell(row=r, column=url_col).value).strip() != ""
        ]
        target_total = len(target_rows) if target_rows else max(1, end_r - start_r + 1)

        processed_count = 0
        logger.info(f"Rentang baris target: {start_r} s.d. {end_r} (Ditemukan {len(target_rows)} baris penugasan)")

        try:
            for row_idx in range(start_r, end_r + 1):
                # Periksa sinyal stop
                if self.event_bus.is_stop_requested():
                    logger.warning("Proses dihentikan oleh pengguna.")
                    break

                # Periksa sinyal jeda
                self.event_bus.wait_if_paused()

                url_val = sheet.cell(row=row_idx, column=url_col).value if url_col else None
                if not url_val or str(url_val).strip() == "":
                    continue

                processed_count += 1
                assignment_id = str(sheet.cell(row=row_idx, column=id_col).value or "").strip()
                nama_usaha = str(sheet.cell(row=row_idx, column=name_col).value or "").strip()
                status_val = str(sheet.cell(row=row_idx, column=status_col).value or "").strip()

                # Emit progress update
                self.event_bus.emit("progress_updated", {
                    "current": processed_count,
                    "total": target_total,
                    "row_idx": row_idx,
                    "nama_usaha": nama_usaha,
                    "assignment_id": assignment_id
                })

                if status_val.upper() in ["SUKSES", "DONE"]:
                    existing_ket = str(sheet.cell(row=row_idx, column=ket_col).value or "").strip() if ket_col else ""
                    logger.info(f"Baris {row_idx} [{nama_usaha}] berstatus {status_val} -> Dilewati (keterangan dipertahankan).")
                    skipped_count += 1
                    self.event_bus.emit("row_completed", {
                        "row_idx": row_idx,
                        "status": f"{status_val} (LEWAT)",
                        "keterangan": existing_ket
                    })
                    continue

                logger.bot(f"--- [Baris {row_idx}/{end_r}] Memproses: {nama_usaha} ---")
                self.event_bus.emit("row_started", {
                    "row_idx": row_idx,
                    "nama_usaha": nama_usaha,
                    "id": assignment_id
                })

                audit_data = {
                    "row_index": row_idx,
                    "assignment_id": assignment_id,
                    "nama_usaha": nama_usaha,
                    "link_fasih": str(url_val).strip(),
                    "status": "GAGAL"
                }

                try:
                    # STEP 1: Buka URL Penugasan (dengan normalisasi protokol https://)
                    raw_url = str(url_val).strip()
                    if not raw_url.startswith(("http://", "https://")):
                        raw_url = "https://" + raw_url

                    nav_loaded = False
                    for nav_attempt in range(3):
                        try:
                            page.wait_for_load_state("domcontentloaded", timeout=4000)
                        except Exception:
                            pass

                        try:
                            page.goto(raw_url, timeout=45000, wait_until="domcontentloaded")
                            nav_loaded = True
                            break
                        except Exception as nav_e:
                            nav_err_msg = str(nav_e).lower()
                            if ("interrupted" in nav_err_msg or "navigation" in nav_err_msg) and nav_attempt < 2:
                                logger.warning(f"Navigasi baris {row_idx} terinterupsi pengalihan browser ({nav_e}). Mencoba ulang dalam 2 detik (Percobaan {nav_attempt+1}/3)...")
                                time.sleep(2.0)
                            else:
                                raise nav_e

                    if not nav_loaded:
                        raise Exception("Gagal memuat URL penugasan setelah 3 percobaan.")

                    random_sleep(1.5, 2.5, event_bus=self.event_bus)

                    # Cek sesi SSO di tengah jalan
                    if BrowserManager.is_session_expired(page) or not BrowserManager.is_page_logged_in(page):
                        page = self._handle_relogin_needed(page, context, row_idx=row_idx, nama_usaha=nama_usaha)
                        if not page or self.event_bus.is_stop_requested():
                            audit_data["status"] = "DIHENTIKAN"
                            audit_data["keterangan"] = "Dihentikan oleh pengguna saat menunggu login SSO"
                            self.event_bus.emit("row_completed", {
                                "row_idx": row_idx,
                                "status": "⏹️ DIHENTIKAN",
                                "keterangan": "Dihentikan oleh pengguna saat menunggu login SSO"
                            })
                            break

                        # KRUSIAL: Muat ulang URL penugasan baris ini karena setelah login browser berada di halaman redirect SSO
                        logger.info(f"🔄 Memuat ulang formulir penugasan Baris {row_idx} [{nama_usaha}]...")
                        page.goto(raw_url, timeout=45000, wait_until="domcontentloaded")
                        random_sleep(1.5, 2.5, event_bus=self.event_bus)

                    # STEP 2: Buka Sub Sidebar Blok II / Nama Usaha
                    if not self.check_pause_and_stop(row_idx, audit_data):
                        break

                    try:
                        page.locator(BLOK_II_XPATH).first.wait_for(state="visible", timeout=12000)
                        safe_click(page, BLOK_II_XPATH, timeout=8000, event_bus=self.event_bus)
                        random_sleep(0.8, 1.4, event_bus=self.event_bus)
                    except Exception as blok2_err:
                        # Jika elemen Blok II tidak muncul, periksa apakah sesi Fasih tiba-tiba habis
                        if BrowserManager.is_session_expired(page) or not BrowserManager.is_page_logged_in(page):
                            page = self._handle_relogin_needed(page, context, row_idx=row_idx, nama_usaha=nama_usaha)
                            if not page or self.event_bus.is_stop_requested():
                                audit_data["status"] = "DIHENTIKAN"
                                audit_data["keterangan"] = "Dihentikan oleh pengguna saat menunggu login SSO"
                                self.event_bus.emit("row_completed", {
                                    "row_idx": row_idx,
                                    "status": "⏹️ DIHENTIKAN",
                                    "keterangan": "Dihentikan oleh pengguna saat menunggu login SSO"
                                })
                                break
                            logger.info(f"🔄 Memuat ulang formulir penugasan Baris {row_idx} [{nama_usaha}] setelah re-login...")
                            page.goto(raw_url, timeout=45000, wait_until="domcontentloaded")
                            random_sleep(1.5, 2.5, event_bus=self.event_bus)
                            try:
                                page.locator(BLOK_II_XPATH).first.wait_for(state="visible", timeout=12000)
                                safe_click(page, BLOK_II_XPATH, timeout=8000, event_bus=self.event_bus)
                                random_sleep(0.8, 1.4, event_bus=self.event_bus)
                            except Exception:
                                pass

                    # Cari item nama usaha di sub-sidebar
                    usaha_item = None
                    if nama_usaha:
                        usaha_xpath = f"xpath=//div[contains(@class, 'fasih-form-sidebar')]//div[contains(@class, 'cursor-pointer') and contains(., '{nama_usaha}')]"
                        if page.locator(usaha_xpath).first.count() > 0:
                            usaha_item = page.locator(usaha_xpath).first

                    if not usaha_item:
                        first_sub = page.locator("xpath=//div[contains(@class, 'fasih-form-sidebar')]//div[contains(@class, 'cursor-pointer') and not(contains(., 'BLOK II')) and not(contains(., 'CATATAN'))]").first
                        if first_sub.count() > 0:
                            usaha_item = first_sub

                    if usaha_item:
                        safe_click(page, usaha_item, timeout=8000, event_bus=self.event_bus)
                        random_sleep(1.5, 2.5, event_bus=self.event_bus)

                    # Tunggu form utama tampil
                    try:
                        page.locator(TOTAL_PENGELUARAN_SELECTOR).first.wait_for(state="attached", timeout=15000)
                    except Exception:
                        if BrowserManager.is_session_expired(page):
                            raise Exception("Sesi Fasih kedaluwarsa saat memuat rincian usaha.")
                        raise Exception("Elemen rincian usaha tidak muncul setelah navigasi.")

                    # STEP 3: Pengeluaran
                    if not self.check_pause_and_stop(row_idx, audit_data):
                        break

                    tp_val = sheet.cell(row=row_idx, column=tp_col).value if tp_col else None
                    tp_raw_val = sheet.cell(row=row_idx, column=tp_raw_col).value if tp_raw_col else None
                    diff_pengeluaran = is_values_different(tp_val, tp_raw_val)

                    audit_data["pengeluaran_clean"] = tp_val
                    audit_data["pengeluaran_raw"] = tp_raw_val

                    if diff_pengeluaran:
                        logger.bot(f"Total Pengeluaran BERBEDA (Clean: {tp_val} vs RAW: {tp_raw_val}) -> Mengisi komponen...")
                        fill_components_from_excel(page, PENGELUARAN_MAP, sheet, row_idx, col_map, event_bus=self.event_bus)
                        remarks_pengeluaran = "#4"
                    else:
                        logger.bot(f"Total Pengeluaran SAMA (Clean: {tp_val} vs RAW: {tp_raw_val}) -> Lewati komponen.")
                        remarks_pengeluaran = "#0"

                    logger.remarks(f"Memberikan remarks '{remarks_pengeluaran}' pada Total Pengeluaran...")
                    add_remarks_to_element(page, TOTAL_PENGELUARAN_SELECTOR, remarks_pengeluaran, event_bus=self.event_bus)
                    audit_data["remarks_pengeluaran"] = remarks_pengeluaran

                    # STEP 4: Pendapatan
                    if not self.check_pause_and_stop(row_idx, audit_data):
                        break

                    tpend_val = sheet.cell(row=row_idx, column=tpend_col).value if tpend_col else None
                    tpend_raw_val = sheet.cell(row=row_idx, column=tpend_raw_col).value if tpend_raw_col else None
                    diff_pendapatan = is_values_different(tpend_val, tpend_raw_val)

                    audit_data["pendapatan_clean"] = tpend_val
                    audit_data["pendapatan_raw"] = tpend_raw_val

                    if diff_pendapatan:
                        logger.bot(f"Total Pendapatan BERBEDA (Clean: {tpend_val} vs RAW: {tpend_raw_val}) -> Mengisi komponen...")
                        fill_components_from_excel(page, PENDAPATAN_MAP, sheet, row_idx, col_map, event_bus=self.event_bus)
                        remarks_pendapatan = "#4"
                    else:
                        logger.bot(f"Total Pendapatan SAMA (Clean: {tpend_val} vs RAW: {tpend_raw_val}) -> Lewati komponen.")
                        remarks_pendapatan = "#0"

                    logger.remarks(f"Memberikan remarks '{remarks_pendapatan}' pada Total Pendapatan...")
                    add_remarks_to_element(page, TOTAL_PENDAPATAN_SELECTOR, remarks_pendapatan, event_bus=self.event_bus)
                    audit_data["remarks_pendapatan"] = remarks_pendapatan

                    # STEP 5: Total Aset
                    if not self.check_pause_and_stop(row_idx, audit_data):
                        break

                    ta_val = sheet.cell(row=row_idx, column=ta_col).value if ta_col else None
                    ta_raw_val = sheet.cell(row=row_idx, column=ta_raw_col).value if ta_raw_col else None
                    diff_aset = is_values_different(ta_val, ta_raw_val)

                    audit_data["aset_clean"] = ta_val
                    audit_data["aset_raw"] = ta_raw_val

                    if diff_aset:
                        tot_aset_num = to_int_or_zero(ta_val)
                        logger.bot(f"Total Aset BERBEDA (Clean: {ta_val} vs RAW: {ta_raw_val}) -> Menghitung target Aset ({tot_aset_num:,})...")

                        elem_tanah = page.locator(ASET_MAP["tanah_bangunan"]["selector"]).first
                        elem_non_tanah = page.locator(ASET_MAP["selain_tanah"]["selector"]).first

                        val_exist_tanah = 0
                        val_exist_lain = 0
                        try:
                            if elem_tanah.count() > 0:
                                inp = elem_tanah.locator("input").first
                                if inp.count() > 0:
                                    val_exist_tanah = parse_fasih_currency(inp.input_value())
                        except Exception:
                            pass

                        try:
                            if elem_non_tanah.count() > 0:
                                inp = elem_non_tanah.locator("input").first
                                if inp.count() > 0:
                                    val_exist_lain = parse_fasih_currency(inp.input_value())
                        except Exception:
                            pass

                        total_exist = val_exist_tanah + val_exist_lain
                        if total_exist > 0:
                            prop_tanah = val_exist_tanah / float(total_exist)
                            aset_tanah_baru = int(round(tot_aset_num * prop_tanah))
                            aset_non_tanah_baru = tot_aset_num - aset_tanah_baru
                        else:
                            pct_tanah = random.randint(60, 70)
                            aset_tanah_baru = int(round(tot_aset_num * (pct_tanah / 100.0)))
                            aset_non_tanah_baru = tot_aset_num - aset_tanah_baru

                        audit_data["aset_tanah_baru"] = aset_tanah_baru
                        audit_data["aset_non_tanah_baru"] = aset_non_tanah_baru

                        if elem_tanah.count() > 0 and elem_tanah.is_visible():
                            safe_fill_currency(page, elem_tanah, aset_tanah_baru, event_bus=self.event_bus)
                            random_sleep(0.3, 0.6, event_bus=self.event_bus, allow_hesitation=False)

                        if elem_non_tanah.count() > 0 and elem_non_tanah.is_visible():
                            safe_fill_currency(page, elem_non_tanah, aset_non_tanah_baru, event_bus=self.event_bus)
                            random_sleep(0.3, 0.6, event_bus=self.event_bus, allow_hesitation=False)

                        remarks_aset = "#4"
                    else:
                        logger.bot(f"Total Aset SAMA (Clean: {ta_val} vs RAW: {ta_raw_val}) -> Lewati pengisian aset.")
                        remarks_aset = "#0"

                    logger.remarks(f"Memberikan remarks '{remarks_aset}' pada Total Aset...")
                    try:
                        add_remarks_to_element(page, ASET_MAP["total_aset"]["selector"], remarks_aset, event_bus=self.event_bus)
                    except Exception:
                        add_remarks_to_element(page, ASET_MAP["tanah_bangunan"]["selector"], remarks_aset, event_bus=self.event_bus)
                    audit_data["remarks_aset"] = remarks_aset

                    # STEP 6: Menuju Blok Catatan di Sidebar
                    if not self.check_pause_and_stop(row_idx, audit_data):
                        break

                    try:
                        page.locator(BLOK_CATATAN_XPATH).first.wait_for(state="visible", timeout=12000)
                        safe_click(page, BLOK_CATATAN_XPATH, timeout=8000, event_bus=self.event_bus)
                    except Exception:
                        safe_click(page, "div[title='CATATAN']", timeout=8000, event_bus=self.event_bus)
                    random_sleep(1.2, 2.0, event_bus=self.event_bus)

                    # STEP 7: Switch Tampilkan Anomali Usaha dan Keluarga
                    if not self.check_pause_and_stop(row_idx, audit_data):
                        break

                    # Pastikan switch aktif: kalau belum -> diaktifkan, kalau sudah -> dibiarkan.
                    # Gagal mengaktifkan = baris GAGAL (tidak boleh submit tanpa anomali ditampilkan).
                    ensure_anomali_switch_active(page, event_bus=self.event_bus)

                    # STEP 8: Cek Galat & Submit
                    if not self.check_pause_and_stop(row_idx, audit_data):
                        break

                    galat_count = 0
                    try:
                        galat_badge = page.locator("xpath=//button[contains(., 'Galat') or contains(., 'galat')]").first
                        if galat_badge.count() > 0 and galat_badge.is_visible():
                            txt = galat_badge.inner_text()
                            nums = re.findall(r"\d+", txt)
                            if nums:
                                galat_count = int(nums[0])
                    except Exception:
                        pass

                    audit_data["galat_count"] = galat_count

                    if galat_count > 0:
                        err_msg = f"Terdapat {galat_count} Galat pada form. Submit dibatalkan secara aman."
                        logger.warning(f"Baris {row_idx} [{nama_usaha}]: {err_msg}")
                        sheet.cell(row=row_idx, column=status_col).value = f"GAGAL ({galat_count} Galat)"
                        sheet.cell(row=row_idx, column=ket_col).value = err_msg
                        audit_data["status"] = "GAGAL"
                        audit_data["keterangan"] = err_msg
                        error_count += 1
                    else:
                        # ------------------------------------------------------------------
                        # STEP 8 SUBMIT FLOW BERANTAI (4 LANGKAH):
                        # 1. Klik Kirim pada formulir Fasih
                        # 2. Muncul dialog pop up ringkasan galat, klik Kirim lagi
                        # 3. Muncul dialog pop up konfirmasi kirim, klik Konfirmasi
                        # 4. Loading dan menunggu dialog 'Assignment Berhasil Disubmit'
                        # ------------------------------------------------------------------

                        # LANGKAH 1: Klik Kirim pada formulir Fasih
                        logger.bot("Galat = 0. [1/4] Mengklik tombol Kirim formulir Fasih...")
                        form_kirim_btn = page.locator("button#btn-kirim, button:has-text('Kirim')").first
                        if form_kirim_btn.count() == 0 or not form_kirim_btn.is_visible():
                            form_kirim_btn = page.locator(KIRIM_BUTTON_XPATH).first

                        if form_kirim_btn.count() == 0:
                            raise Exception("Tombol Kirim pada formulir Fasih tidak ditemukan.")

                        safe_click(page, form_kirim_btn, timeout=8000, event_bus=self.event_bus, force=True)
                        # Beri jeda render animasi agar pop up ringkasan galat benar-benar muncul di layar
                        random_sleep(1.8, 2.5, event_bus=self.event_bus, allow_hesitation=False)

                        # LANGKAH 2: Dialog Pop-up Ringkasan Galat -> Klik Kirim Lagi
                        btn_konfirmasi = page.locator(KONFIRMASI_KIRIM_BUTTON_SELECTOR).last

                        # Jika tombol konfirmasi belum langsung ada, berarti dialog ringkasan galat sedang terbuka
                        if not btn_konfirmasi.is_visible(timeout=2000):
                            logger.bot("[2/4] Menunggu pop-up Ringkasan Galat...")
                            ringkasan_kirim_btn = page.locator(RINGKASAN_KIRIM_SELECTOR).last

                            try:
                                ringkasan_kirim_btn.wait_for(state="visible", timeout=12000)
                                random_sleep(0.5, 1.0, event_bus=self.event_bus, allow_hesitation=False)
                                if ringkasan_kirim_btn.is_disabled():
                                    raise Exception("Tombol Kirim pada dialog Ringkasan Galat tidak aktif (terdapat galat yang belum diselesaikan).")
                                logger.bot("[2/4] Dialog Ringkasan Galat muncul. Mengklik tombol Kirim...")
                                safe_click(page, ringkasan_kirim_btn, timeout=8000, event_bus=self.event_bus, force=True)
                                random_sleep(1.5, 2.2, event_bus=self.event_bus, allow_hesitation=False)
                            except Exception as e_ringkasan:
                                # Jika tombol konfirmasi ternyata sudah muncul tanpa perlu klik ringkasan
                                if not btn_konfirmasi.is_visible(timeout=2000):
                                    raise Exception(f"Gagal memproses dialog Ringkasan Galat: {e_ringkasan}")

                        # LANGKAH 3: Dialog Pop-up Konfirmasi Kirim -> Klik Konfirmasi
                        logger.bot("[3/4] Menunggu pop-up Konfirmasi Kirim...")
                        btn_konfirmasi = page.locator(KONFIRMASI_KIRIM_BUTTON_SELECTOR).last
                        btn_konfirmasi.wait_for(state="visible", timeout=12000)
                        random_sleep(0.5, 1.0, event_bus=self.event_bus, allow_hesitation=False)

                        if self.debug_review_submit:
                            logger.warning(f"🔍 [MODE REVIEW DEBUGGING] Baris {row_idx} [{nama_usaha}]: Modal Konfirmasi Kirim terbuka.")
                            logger.info("Silakan periksa isian data di browser. Tekan tombol '✅ Konfirmasi Submit & Lanjut' di aplikasi untuk mengirimkan.")
                            WindowController.restore_and_focus_browser()
                            self.event_bus.emit("waiting_confirm_submit", {
                                "row_idx": row_idx,
                                "nama_usaha": nama_usaha
                            })
                            self.event_bus.wait_for_submit_confirmation()

                            if self.event_bus.is_stop_requested():
                                logger.warning("Proses dihentikan sebelum konfirmasi submit.")
                                audit_data["status"] = "DIHENTIKAN"
                                audit_data["keterangan"] = "Dihentikan oleh pengguna saat review konfirmasi"
                                self.event_bus.emit("row_completed", {
                                    "row_idx": row_idx,
                                    "status": "⏹️ DIHENTIKAN",
                                    "keterangan": "Dihentikan oleh pengguna saat review konfirmasi"
                                })
                                break

                            logger.bot("Konfirmasi diterima dari pengguna. Mengklik tombol Konfirmasi...")
                            safe_click(page, btn_konfirmasi, timeout=8000, event_bus=self.event_bus, force=True)
                            if self.auto_minimize:
                                WindowController.minimize_browser()
                        else:
                            # Mode Automasi Beneran: Otomatis klik tombol Konfirmasi
                            logger.bot("[3/4] Dialog Konfirmasi Kirim aktif. Mengklik tombol Konfirmasi...")
                            safe_click(page, btn_konfirmasi, timeout=8000, event_bus=self.event_bus, force=True)

                        # LANGKAH 4: Menunggu Dialog 'Assignment Berhasil Disubmit'
                        logger.bot("[4/4] Menunggu server memproses submit (dialog 'Assignment Berhasil Disubmit')...")
                        success_dialog = page.locator(SUCCESS_SUBMIT_DIALOG_SELECTOR).first

                        try:
                            success_dialog.wait_for(state="visible", timeout=35000)
                            logger.success(f"🎉 Baris {row_idx} [{nama_usaha}]: Dialog 'Assignment Berhasil Disubmit' terkonfirmasi!")

                            # Tutup modal secara bersih jika tombol 'Tetap di Halaman' tersedia
                            try:
                                btn_tetap = page.locator("button:has-text('Tetap di Halaman')").first
                                if btn_tetap.is_visible(timeout=1500):
                                    safe_click(page, btn_tetap, timeout=3000, event_bus=self.event_bus)
                            except Exception:
                                pass

                            sheet.cell(row=row_idx, column=status_col).value = "SUKSES"
                            # sheet.cell(row=row_idx, column=ket_col).value = f"Assignment Berhasil Disubmit ({datetime.now().strftime('%Y-%m-%d %H:%M:%S')})"
                            sheet.cell(row=row_idx, column=ket_col).value = "Assignment Berhasil Disubmit"
                            audit_data["status"] = "SUKSES"
                            audit_data["keterangan"] = "Assignment Berhasil Disubmit"
                            done_session += 1
                            logger.success(f"✅ Baris {row_idx} [{nama_usaha}] BERHASIL DISUBMIT (SUKSES).")

                        except Exception as wait_err:
                            # Periksa apakah ada notifikasi error/galat dari server
                            err_dialog = page.locator("div[role='alertdialog']:has-text('Gagal'), div:has-text('Gagal mengirimkan'), div:has-text('Terjadi kesalahan')").first
                            if err_dialog.count() > 0 and err_dialog.is_visible():
                                err_text = err_dialog.inner_text().strip().replace("\n", " ")
                                raise Exception(f"Submit ditolak server: {err_text[:120]}")
                            raise Exception(f"Timeout menunggu dialog 'Assignment Berhasil Disubmit' (server belum merespons dalam 35 detik): {wait_err}")

                except Exception as row_err:
                    error_msg = str(row_err)
                    logger.error(f"❌ Baris {row_idx} [{nama_usaha}] GAGAL: {error_msg}")
                    sheet.cell(row=row_idx, column=status_col).value = "GAGAL"
                    sheet.cell(row=row_idx, column=ket_col).value = error_msg[:150]
                    audit_data["status"] = "GAGAL"
                    audit_data["keterangan"] = error_msg
                    error_count += 1

                # Simpan checkpoint ke Excel secara aman per baris
                try:
                    wb.save(self.excel_path)
                except Exception as save_err:
                    logger.warning(f"Gagal menyimpan checkpoint ke Excel (pastikan file tidak dibuka di Excel): {save_err}")

                # Catat ke audit log
                audit_logger.record_row(audit_data)

                # Emit row completed
                self.event_bus.emit("row_completed", {
                    "row_idx": row_idx,
                    "status": audit_data["status"],
                    "keterangan": audit_data["keterangan"]
                })

                random_sleep(1.0, 2.0, event_bus=self.event_bus)

        finally:
            if wb:
                try:
                    wb.save(self.excel_path)
                    wb.close()
                except Exception:
                    pass
            if browser:
                try:
                    browser.close()
                except Exception:
                    pass
            if p:
                try:
                    p.stop()
                except Exception:
                    pass

            # Restore browser jika sebelumnya di-minimize
            if self.auto_minimize:
                try:
                    WindowController.restore_and_focus_browser()
                except Exception:
                    pass

            was_stopped = self.event_bus.is_stop_requested()
            self.event_bus.set_state(EngineState.IDLE)
            if was_stopped:
                logger.warning(f"⏹️ Otomasi Dihentikan oleh Pengguna! (Sukses Sesi Ini: {done_session}, Gagal: {error_count}, Dilewati: {skipped_count})")
            else:
                logger.success(f"🏁 Eksekusi Data Cleaning Selesai! (Sukses Sesi Ini: {done_session}, Gagal: {error_count}, Dilewati: {skipped_count})")

            self.event_bus.emit("all_completed", {
                "done": done_session,
                "error": error_count,
                "skipped": skipped_count,
                "was_stopped": was_stopped
            })
